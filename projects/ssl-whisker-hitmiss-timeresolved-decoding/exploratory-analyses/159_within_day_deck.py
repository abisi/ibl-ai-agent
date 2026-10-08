"""159 -- Within-day deck (user 2026-10-08: "like for the prelick convergence, build ppt slides summarizing the results, focusing on the
within-day change"; expert data hidden). Methods first (derived step by step, with equations and schematics), then whole-brain results
(active trials: is choice decoding stable within the session; passive pre vs post: does the whisker response move on the hit / miss axis,
beyond session time), then per area group; summary, caveats, backup figures. The pre-lick results of the sister project
(ssl-prelick-convergence) are not repeated here.
Figures and numbers come from the 155 helpers (re-plotted from the per-session tables) and the 158 COSYNE figure.
Output: combined_results_ks4/<slug>/report/within_day_deck.pptx. Run (haas, repo root): python .../159_within_day_deck.py
"""
from __future__ import annotations

import importlib
import json
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

EA = Path(__file__).resolve().parent
sys.path.insert(0, str(EA)); sys.path.insert(0, str(EA.parents[2] / "scripts"))
D = importlib.import_module("155_deck_pptx")
H, COL, COH, P, PUB = D.H, D.COL, D.COH, D.P, D.PUB
OUT = D.RDIR / "within_day_deck.pptx"
D.SEC_TAG.update({"M": "Methods", "A": "Whole brain, active trials: choice decoding within the session",
                  "B": "Whole brain, passive pre vs post: the whisker response on the hit / miss axis", "C": "Per area group",
                  "S": "Summary", "bk": "Backup"})
new, text, bullets, picture, takehome, eqbox = D.new_slide, D.text, D.bullets, D.picture, D.takehome, D.eqbox


def fig_windows():
    """trial windows: baseline -55..-20 ms, response 5-35 ms after stimulus onset; active trials with a lick before 35 ms excluded"""
    fig, ax = plt.subplots(figsize=(3.6, 1.3)); ax.set_xlim(-80, 120); ax.set_ylim(0, 1); ax.set_yticks([])
    for s in ("left", "right", "top"):
        ax.spines[s].set_visible(False)
    ax.axvspan(-55, -20, ymin=0.2, ymax=0.7, color="0.85"); ax.text(-37.5, 0.78, "baseline\n-55 to -20 ms", ha="center", fontsize=5.5)
    ax.axvspan(5, 35, ymin=0.2, ymax=0.7, color="#f4c7a1"); ax.text(20, 0.78, "response\n5 to 35 ms", ha="center", fontsize=5.5)
    ax.axvline(0, color="black", lw=1); ax.text(0, 0.05, "stimulus", ha="center", fontsize=5.5)
    ax.text(80, 0.45, "licks start > 35 ms\n(active trials with an\nearlier lick excluded)", fontsize=5, ha="center")
    ax.set_xlabel("time from stimulus onset (ms)")
    return D.savefig(fig, "w_windows")


def fig_stats():
    fig, ax = plt.subplots(figsize=(4.4, 1.3)); ax.axis("off"); ax.set_xlim(0, 10); ax.set_ylim(0, 3)
    for i, (c, n) in enumerate((("R+", 3), ("R-", 3))):
        for j in range(n):
            x = 0.4 + i * 4.8 + j * 1.4
            ax.add_patch(plt.Rectangle((x, 1.6), 1.1, 0.8, fc=COL[c], alpha=0.25, ec=COL[c]))
            ax.text(x + 0.55, 2.0, f"mouse\n1 session", ha="center", va="center", fontsize=4.8)
    ax.text(5, 0.9, "one learning session per mouse = one value per session (unit of analysis)", ha="center", fontsize=6)
    ax.text(5, 0.3, "within cohort: Wilcoxon (| paired t); R+ vs R-: Mann-Whitney U (| Welch t); uncorrected", ha="center", fontsize=6)
    return D.savefig(fig, "w_stats")


def crops_158():
    """split the one-line 158 figure into its panels"""
    im = Image.open(PUB / "158_cosyne_within_day_learners.png").convert("RGB"); w, h = im.size
    out = {}
    for k, (x0, x1) in {"a": (0.0, 0.175), "b": (0.175, 0.345), "c": (0.345, 0.5), "d": (0.5, 0.665), "e": (0.665, 1.0),
                        "bc": (0.175, 0.5)}.items():
        p = D.OUTF / f"w158_{k}.png"; im.crop((int(x0 * w), 0, int(x1 * w), h)).save(p, dpi=(D.DPI, D.DPI)); out[k] = p
    return out


def main():
    H.setup(); plt.rcParams.update({"mathtext.fontset": "dejavusans"})
    R = {}
    N = json.loads((D.RDIR / "numbers.json").read_text())
    A = D.load_135()
    C = pd.read_parquet(EA / "140_coding_direction_noise.parquet"); C = C[(C.split == "hitmedian") & C.skipped_reason.isna()]
    V = pd.read_parquet(EA / "146_choice_axis_readout.parquet")
    V = V[(V.area == "whole_brain") & (V.unit_set == "stable") & V.skipped_reason.isna()]
    VE, VB, VT = (V[V.response == r] for r in ("epochbase", "baseline", "trialbase"))
    L = H.learners(); AL = A[H.mouse_of(A).isin(L)]
    S158 = pd.read_csv(EA / "158_stats_learners.csv")
    figs = dict(task=D.fig_task(), timeline=D.fig_timeline(), dec=D.fig_decode_schematic(), splits=D.fig_splits(R), lt=D.fig_lt(R),
                ltr=D.fig_lt_robust(R), pa=D.fig_pattern_axis(), null=D.fig_null(), epochs=D.fig_lick_epochs(R),
                b4=D.excess_fig("b4", [(A, "shift_excess_dWR", "whisker\ncosine"), (A, "shift_excess_dAR", "auditory\ncosine"),
                                       (A, "shift_excess_dWAP", "whisker - auditory\nprojection"),
                                       (VT, "shift_excess_dWA", "decoder: whisker -\nauditory readout")], R, w=6.9),
                b5=D.excess_fig("b5", [(AL, "shift_excess_dWR", "learners:\nwhisker cosine"), (AL, "shift_excess_dWAP", "learners:\nW - A projection"),
                                       (C, "shift_excess_daxisR", "late CD: passive\naxis cosine"), (C, "shift_excess_dWAP", "late CD:\nW - A projection")], R, w=7.6),
                state=D.fig_state(R), expo=D.fig_exposure(R), dtrain=D.fig_dec_train(), dro=D.fig_dec_readout(R, VE), dnull=D.fig_dec_null(R, VE, VB),
                bridge=D.fig_bridge(), windows=fig_windows(), stats=fig_stats())
    c158 = crops_158()
    E = dict(
        resp=D.eq("w_resp", [r"$r_{it} = f^{[5,35]}_{it} - \langle f^{[-55,-20]}_{it'} \rangle_{t' \in \mathrm{epoch}}$",
                             r"$\varepsilon_{it} = r_{it}/\sigma_i$   (evoked, scaled by the unit's SD over all trials; not centred)"]),
        dec=D.eq("w_dec", [r"$\Delta\mathrm{acc} = \mathrm{acc}_{\mathrm{real}} - \langle \mathrm{acc}^{(k)}_{\mathrm{shift}} \rangle_k$"]),
        split=D.eq("w_split", [r"midpoint: $t_{n/2}$;   hit-median: hit $\lfloor H/2 \rfloor + 1$"]),
        lt=D.eq("w_lt", [r"$\delta_s = \Delta\mathrm{acc}^{\mathrm{after}}_s - \Delta\mathrm{acc}^{\mathrm{before}}_s$",
                         r"$\mathrm{excess} = \delta_{\mathrm{LT}} - \langle \delta_s \rangle_{s \in \mathrm{placebo}}$"]),
        lick=D.eq("w_lick", [r"$\mathbf{L} = \bar{\mathbf{z}}_{\mathrm{hit}} - \bar{\mathbf{z}}_{\mathrm{miss}}$  (active whisker trials)",
                             r"$\mathbf{p}_{W} = \langle \boldsymbol{\varepsilon}_t \rangle_{t \in \mathrm{whisker,\ epoch}}$,  same for $\mathbf{p}_A$"]),
        cos=D.eq("w_cos", [r"$\cos\theta = \dfrac{\mathbf{p}\cdot\mathbf{L}}{\|\mathbf{p}\|\,\|\mathbf{L}\|},\qquad \pi = \dfrac{\mathbf{p}\cdot\hat{\mathbf{L}}}{\sqrt{n}}$"]),
        ro=D.eq("w_ro", [r"$\mathrm{readout} = \dfrac{s - \frac{1}{2}(\bar{s}_{\mathrm{hit}} + \bar{s}_{\mathrm{miss}})}{\mathrm{SD}(s_{\mathrm{active}})}$"]),
        excess=D.eq("w_excess", [r"$\Delta M = M^{\mathrm{post}} - M^{\mathrm{pre}}$",
                                 r"$\mathrm{excess} = \Delta M - \dfrac{1}{K}\sum_{k=1}^{K}\Delta M^{(k)}_{\mathrm{shift}}$"]),
        contrast=D.eq("w_contrast", [r"$\Delta\pi_W - \Delta\pi_A$   (a drift shared by both stimuli cancels)"]),
        ss=D.eq("w_ss", [r"$\hat{\mathbf{x}} = \mathbf{L}/\|\mathbf{L}\|,\quad \hat{\mathbf{y}} \propto \mathbf{p}^{\mathrm{pre}}_W - (\mathbf{p}^{\mathrm{pre}}_W\cdot\hat{\mathbf{x}})\,\hat{\mathbf{x}}$",
                         r"$\Delta x = (\mathbf{p}^{\mathrm{post}} - \mathbf{p}^{\mathrm{pre}})\cdot\hat{\mathbf{x}}$"]),
        cd2=D.eq("w_cd2", [r"$\mathbf{CD}_2 = \bar{\mathbf{z}}_{\mathrm{hit},2} - \bar{\mathbf{z}}_{\mathrm{miss},2}$  (second half of a hit-median split)"]),
        expo=D.eq("w_expo", [r"$\Delta M = \beta_0 + \beta_{R-}\,\mathbb{1}_{R-} + \beta_W z(n_W) + \beta_A z(n_A) + \epsilon$"]))
    g = lambda pnl, m, c: S158[(S158.panel == pnl) & (S158.measure == m) & (S158.cohort == c)].iloc[0]

    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    D.CUR.update(sec="intro", n=1)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.7, 0.9, 12, 1.4, "Within one learning session, the reward contingency moves the whisker response on the hit / miss axis", size=30, bold=True)
    text(s, 0.7, 2.5, 12, 0.5, "Whole-brain Neuropixels, whisker learning day (day 0); R+ (whisker licks rewarded) vs R- (not rewarded); within-day change only",
         size=16)
    for i, (k, t) in enumerate((("Methods", "Epochs, windows, the hit / miss axis, projections, decoder readout, the session-time null, state space"),
                                ("Whole brain", "Active trials: choice decoding within the session.  Passive pre vs post: the whisker response on the hit / miss axis"),
                                ("Per area", "Area groups with >= 3 sessions per cohort"), ("Summary", "COSYNE figure, synthesis, caveats"))):
        text(s, 0.7, 3.7 + i * 0.8, 2.2, 0.6, k, size=20, bold=True); text(s, 2.9, 3.75 + i * 0.8, 10, 0.7, t, size=16)

    # ------------------------------------------------------------------ methods
    D.section_slide(prs, "M", "Methods", "From spikes to a position on the hit / miss axis, step by step")
    s = new(prs, "Same stimuli, same action, different contingency", "One learning session (whisker day 0) per mouse")
    picture(s, figs["task"], 0.5, 1.4, w=6.2); picture(s, figs["timeline"], 0.5, 5.2, w=6.2)
    bullets(s, 7.2, 1.5, 5.7, 5, ["Hit = lick after the whisker stimulus: same label in both cohorts (trained response in R+, error in R-).",
                                  "Passive pre: whisker and auditory stimuli before the task, no reward; active: whisker, auditory and catch trials; "
                                  "passive post: the same passive stimuli after the task.",
                                  "Question 1 (active): does hit / miss decoding change within the session?",
                                  "Question 2 (passive): after the task, does the whisker response look more like a hit or a miss, and differently in R+ and R-?"],
            size=15)
    s = new(prs, "Step 1: the earliest response, before any lick", "Response 5-35 ms after stimulus onset, minus the unit's baseline in that epoch")
    picture(s, figs["windows"], 0.5, 1.4, w=6.3)
    eqbox(s, 7.1, 1.3, 5.9, [("response of unit i on trial t, and evoked scaling", E["resp"])])
    bullets(s, 0.6, 4.4, 12.2, 2.6, ["Epoch baseline (-55 to -20 ms, averaged over the epoch's trials) removes slow rate changes between epochs; variants: "
                                     "per-trial baseline, and the baseline window alone (state without the sensory response, a control).",
                                     "Units: shared tracked stable units (137b): firing >= 0.5 Hz in passive pre, passive post and both active halves, drift-checked; "
                                     "the same units in every epoch.",
                                     "Sessions: learning day with passive trials before and after the task, >= 3 whisker hits and >= 3 misses (no lick before 35 ms), "
                                     ">= 20 tracked units (whole brain; per area group too)."], size=13)
    s = new(prs, "Step 2: a hit / miss axis from the active trials", "The direction in which population activity differs between whisker trials with and without a lick")
    picture(s, figs["pa"], 0.6, 1.4, w=4.4)
    eqbox(s, 5.6, 1.3, 7.4, [("lick axis (mean difference of z-scored responses) and the mean evoked pattern of a condition", E["lick"]),
                              ("alignment (cosine) and reach (projection on the unit axis, per unit)", E["cos"])])
    bullets(s, 0.6, 4.9, 12.2, 2.0, ["The axis is built from active trials only; passive patterns are measured on it, never used to build it.",
                                     "Projection = size × cosine: it moves if the response grows, shrinks or turns relative to the axis; the cosine isolates the direction.",
                                     "Reliability of the 5-35 ms axis is modest (split-half cosine); tests use raw cosines and projections."], size=13)
    s = new(prs, "Step 3: a choice decoder as a readout", "Trained on active hit vs miss; applied to trials it never saw")
    picture(s, figs["dtrain"], 0.5, 1.4, w=5.2)
    eqbox(s, 6.2, 1.3, 6.8, [("score s of a trial; > 0 hit side, < 0 miss side (SD of the held-out active scores)", E["ro"])])
    bullets(s, 6.2, 3.2, 6.8, 3.4, ["L2 logistic regression; balanced subsamples of hits and misses, stratified k-fold.",
                                    "Held-out active whisker trials, all active auditory and all passive trials are scored by models that did not train on them.",
                                    "Same question as the axis, with a weighted direction instead of a mean difference."], size=13)
    s = new(prs, "Step 4: session time is a confound", "Hit / miss labels drift within the session (R- hits early, misses late); any hit / miss axis can learn 'early vs late'")
    picture(s, figs["null"], 0.5, 1.4, w=7.6)
    eqbox(s, 8.5, 1.3, 4.5, [("change of a metric and its excess over K axes rebuilt from shifted labels", E["excess"])])
    bullets(s, 0.6, 4.9, 12.2, 2.0, ["Linear shift: labels shifted by 10-50 % of the time-ordered active whisker trials, no wrap-around; keeps the drift of labels and activity.",
                                     "Shuffled labels destroy trial order (too lenient); a circular shift turns a trend into a sawtooth.",
                                     "Conservative: genuine learning is time-correlated too."], size=13)
    s = new(prs, "Step 5: an auditory control and a state space", "Auditory licks are rewarded in both cohorts")
    eqbox(s, 0.5, 1.3, 6.0, [("whisker minus auditory change: drift shared by both stimuli cancels", E["contrast"]),
                              ("state space: the hit / miss axis and the passive-pre whisker pattern orthogonal to it", E["ss"])])
    picture(s, c158["a"], 7.0, 1.3, w=5.6, h=4.2, scale=20)
    bullets(s, 0.6, 4.7, 6.2, 2.3, ["Before the task, the whisker response sits between active hits and misses on x, and far along its own pattern on y.",
                                    "Hypothesis: after the task it moves toward hits in R+, toward misses in R-."], size=13)
    s = new(prs, "Step 6: statistics", "Exploratory; uncorrected")
    picture(s, figs["stats"], 0.6, 1.4, w=7.5)
    bullets(s, 0.6, 4.0, 12.2, 3.0, ["Unit of analysis: session (one learning session per mouse).",
                                     "Within cohort: Wilcoxon signed-rank (and paired / one-sample t); R+ vs R-: Mann-Whitney U (and Welch t); shown as p | p.",
                                     "COSYNE figure: non-parametric tests only.  Learners: mice with a good or moderate learning category; main tables use all mice "
                                     "with learners as robustness, the COSYNE figure learners only."], size=14)

    # ------------------------------------------------------------------ A
    D.section_slide(prs, "A", "Active trials: is choice decoding stable within the session?", "Hit / miss decoding against a linear-shift null")
    s = new(prs, "Choice is decodable early after the whisker, equally in both cohorts", "Decode hit vs miss from the 5-50 ms response; subtract a time-preserving null")
    picture(s, figs["dec"], 0.5, 1.4, w=6.4)
    eqbox(s, 0.5, 5.0, 6.4, [("accuracy above the linear-shift null", E["dec"])])
    k0 = lambda k: N.get(k, np.nan)
    bullets(s, 7.3, 1.5, 5.6, 3.5, [f"Session-wide, 5-50 ms: R+ {k0('k0_rp'):+.3f}, R- {k0('k0_rm'):+.3f} above the null.",
                                    f"No cohort difference (p = {H.pnum(k0('k0_mw'))} | {H.pnum(k0('k0_w'))})."], size=15)
    takehome(s, 7.3, 5.0, 5.6, 1.4, "A lick after the whisker is predictable from early whole-brain activity, rewarded or not.")
    sm, sh, hr = R["split_mid"], R["split_hitmedian"], R["split_hr"]
    s = new(prs, "Splitting the session in two: no change in choice decoding", "Midpoint (time) and hit-median (equal hits) splits; 5-100 ms")
    picture(s, figs["splits"], 0.3, 1.3, w=8.6)
    eqbox(s, 9.2, 1.3, 3.8, [("split points", E["split"])])
    bullets(s, 9.2, 2.6, 3.8, 3.0, [f"Midpoint: R+ {sm['R+'][0]:+.3f}, R- {sm['R-'][0]:+.3f}; R+ vs R- p = {P(*sm['cohort'])}.",
                                    f"Hit-median: R+ {sh['R+'][0]:+.3f}, R- {sh['R-'][0]:+.3f}; R+ vs R- p = {P(*sh['cohort'])}.",
                                    f"Hit rate changes between halves (R+ p = {P(hr['R+'][1], hr['R+'][2])}, R- p = {P(hr['R-'][1], hr['R-'][2])})."], size=13)
    takehome(s, 9.2, 5.6, 3.8, 1.3, "Decodability is stable between halves in both cohorts.")
    s = new(prs, "At the learning trial, R+ decoding rises beyond placebo splits", "Split at the behavioural learning trial; control: every other split of the same session")
    picture(s, figs["lt"], 0.3, 1.3, w=8.6)
    eqbox(s, 9.2, 1.3, 3.8, [("change at a split s and excess over placebo splits", E["lt"])])
    bullets(s, 9.2, 3.25, 3.8, 2.4, [f"R+ excess {R['lt_ex_R+'][0]:+.3f} (p = {P(R['lt_ex_R+'][1], R['lt_ex_R+'][2])}); R- {R['lt_ex_R-'][0]:+.3f}.",
                                     f"R+ vs R- p = {P(R['lt_ex_cohort'][0], R['lt_ex_cohort'][1])}; permutation p = {H.pnum(R['lt_ex_cohort'][2])}."], size=13)
    takehome(s, 9.2, 5.6, 3.8, 1.3, "A small R+-specific gain appears when the split is tied to behaviour.")
    s = new(prs, "The learning-trial gain holds across windows and belongs to abrupt learners", "L5: Bayesian change point of the whisker hit sequence")
    picture(s, figs["ltr"], 0.4, 1.3, w=8.4)
    bl = [f"{w[0].replace('ms', ' ms')}: R+ {R[f'l5_{w[0]}_R+'][0]:+.3f} (p = {P(R[f'l5_{w[0]}_R+'][1], R[f'l5_{w[0]}_R+'][2])}); R+ vs R- p = "
          f"{P(*R[f'l5_{w[0]}_cohort'])}" for w in H.WINS]
    bullets(s, 9.2, 1.4, 3.8, 3.5, bl + [f"R+ step ({R['sg'][2]}) vs gradual ({R['sg'][3]}) learners: p = {P(R['sg'][0], R['sg'][1])}."], size=12)
    takehome(s, 9.2, 5.4, 3.8, 1.5, "Decoding accuracy changes little within the session; the effect is small and tied to abrupt learning.")

    # ------------------------------------------------------------------ B
    D.section_slide(prs, "B", "Passive pre vs post: the whisker response on the hit / miss axis", "Does the earliest whisker response move toward hits (R+) or misses (R-)?")
    s = new(prs, "Gain or direction?", "A change in size scales the response along itself; the active trials give a reference direction")
    picture(s, figs["bridge"], 0.4, 1.3, w=7.2)
    bullets(s, 7.9, 1.4, 5.1, 4, ["Passive trials are measured without the task: a shift along the hit / miss axis is a change in the sensory representation itself.",
                                  "Shrinkage along the response's own pattern is shared by both cohorts and scales with trial numbers (exposure slide)."], size=13)
    bpp, cpp = g("b", "W projection post - pre", "R+ vs R-"), g("c", "A projection post - pre", "R+ vs R-")
    bp, bm = g("b", "W projection post - pre", "R+"), g("b", "W projection post - pre", "R-")
    s = new(prs, "After the task, the R- whisker response drops on the hit / miss axis", "Projection of the evoked patterns on the active lick axis, four epochs (learners)")
    picture(s, c158["bc"], 0.4, 1.3, w=8.4, h=4.6, scale=20)
    bullets(s, 9.0, 1.5, 4.0, 4.0, [f"Whisker, post - pre: R+ {bp.mean_a:+.3f} (p = {H.pnum(bp.p_nonparam)}), R- {bm.mean_a:+.3f} (p = {H.pnum(bm.p_nonparam)}); "
                                    f"R+ vs R- p = {H.pnum(bpp.p_nonparam)} (Mann-Whitney).",
                                    f"Auditory (control): R+ vs R- p = {H.pnum(cpp.p_nonparam)}.",
                                    "The R- drop starts in the second active half, when R- mice stop licking to the whisker."], size=13)
    takehome(s, 9.0, 5.5, 4.0, 1.3, "R- whisker responses lose their hit-side component during the task and keep it lost.")
    d = R["d"].iloc[0]
    s = new(prs, "The direction changes, not only the size", "Cosine of the evoked patterns with the lick axis (all mice)")
    picture(s, figs["epochs"], 0.6, 1.3, w=7.6)
    bullets(s, 8.6, 1.5, 4.4, 3.0, [f"Passive post - pre, R+ vs R-: p = {P(d.p_nonparam, d.p_param)}.", "Auditory pattern (dashed): no cohort difference."], size=14)
    takehome(s, 8.6, 5.2, 4.4, 1.5, "The R- whisker response turns away from the lick axis.")
    dro, droA = R["dro_W"], R["dro_A"]
    m_ = lambda d_, c, col: float(np.nanmean(d_[d_.reward_group == c][col]))
    s = new(prs, "A choice decoder reads the R- whisker response as a no-lick trial", "Decoder trained on active hit vs miss (5-35 ms), applied to passive trials")
    picture(s, figs["dtrain"], 0.4, 1.35, w=4.6); picture(s, figs["dro"], 5.6, 1.3, w=7.4, h=3.3)
    bullets(s, 5.6, 4.7, 7.4, 1.4, [f"Whisker: R- {m_(VE, 'R-', 'ro_std_passive_pre_W'):+.2f} -> {m_(VE, 'R-', 'ro_std_passive_post_W'):+.2f} SD "
                                    f"(p = {P(dro['R-'][1], dro['R-'][2])}); R+ {m_(VE, 'R+', 'ro_std_passive_pre_W'):+.2f} -> {m_(VE, 'R+', 'ro_std_passive_post_W'):+.2f}; "
                                    f"R+ vs R- p = {P(*dro['cohort'])}.", f"Auditory: R+ vs R- p = {P(*droA['cohort'])}."], size=12)
    dbs, dcov = R["dbase"], R["dcov"]
    s = new(prs, "But a decoder trained on time-shifted labels does it too", "Shifted labels keep the session's drift; the baseline window carries no sensory response")
    picture(s, figs["dnull"], 0.3, 1.3, w=8.9)
    bullets(s, 9.4, 1.4, 3.6, 3.9, [f"Baseline window alone: R- change {dbs['R-'][0]:+.2f} (p = {P(dbs['R-'][1], dbs['R-'][2])}); R+ vs R- p = {P(*dbs['cohort'])}.",
                                    f"R- coefficient: p = {H.pnum(dcov[0][2])} alone, {H.pnum(dcov[1][2])} with baseline rate and gap, {H.pnum(dcov[2][2])} with reward rate."],
            size=12)
    takehome(s, 9.4, 5.3, 3.6, 1.6, "A decoder can read session time and state: the linear-shift null is needed.")
    dp, dap, dwap = (g("d", m, "R+ vs R-") for m in ("shift_excess_dWP", "shift_excess_dAP", "shift_excess_dWAP"))
    dm_ = g("d", "shift_excess_dWP", "R-")
    s = new(prs, "Beyond session time, the R- whisker response still drops on the hit / miss axis", "Pre -> post projection change minus the change on axes rebuilt from shifted labels (learners)")
    picture(s, c158["d"], 0.6, 1.3, w=5.0, h=5.0, scale=20)
    eqbox(s, 6.0, 1.3, 7.0, [("excess over the linear-shift null", E["excess"]), ("whisker - auditory contrast", E["contrast"])])
    bullets(s, 6.0, 3.9, 7.0, 2.0, [f"Whisker: R- {dm_.mean_a:+.3f} (p = {H.pnum(dm_.p_nonparam)}); R+ vs R- p = {H.pnum(dp.p_nonparam)}.",
                                    f"Auditory: R+ vs R- p = {H.pnum(dap.p_nonparam)}; whisker - auditory: p = {H.pnum(dwap.p_nonparam)} (Mann-Whitney)."], size=13)
    takehome(s, 6.0, 5.6, 7.0, 1.2, "The R- change is not session time and is specific to the whisker response.")
    b4 = R["b4"]; g4 = lambda m: b4[(b4.measure == m) & (b4.cohort == "R+ vs R-")].iloc[0]
    s = new(prs, "Same answer with cosines and with the decoder relative to auditory", "All mice; lick axis (135) and decoder (146) against the linear-shift null")
    picture(s, figs["b4"], 0.3, 1.3, w=8.6)
    bullets(s, 9.2, 1.5, 3.8, 4.5, [f"Whisker cosine p = {P(g4('shift_excess_dWR').p_nonparam, g4('shift_excess_dWR').p_param)}.",
                                    f"Auditory cosine p = {P(g4('shift_excess_dAR').p_nonparam, g4('shift_excess_dAR').p_param)}.",
                                    f"Whisker - auditory projection p = {P(g4('shift_excess_dWAP').p_nonparam, g4('shift_excess_dWAP').p_param)}.",
                                    f"Decoder whisker - auditory p = {P(g4('shift_excess_dWA').p_nonparam, g4('shift_excess_dWA').p_param)}; whisker alone n.s."], size=12)
    b5 = R["b5"]; g5 = lambda i: b5[b5.cohort == "R+ vs R-"].iloc[i]
    s = new(prs, "Robust in learners and with a second hit / miss axis", "Late coding direction of a hit-median split (140)")
    picture(s, figs["b5"], 0.3, 1.3, w=9.4)
    eqbox(s, 9.9, 1.3, 3.1, [("late coding direction", E["cd2"])])
    bullets(s, 9.9, 2.6, 3.1, 3.2, [f"Learners: p = {P(g5(0).p_nonparam, g5(0).p_param)} and {P(g5(1).p_nonparam, g5(1).p_param)}.",
                                    f"Late CD: p = {P(g5(2).p_nonparam, g5(2).p_param)} and {P(g5(3).p_nonparam, g5(3).p_param)}."], size=12)
    ex_ = g("e", "whisker x displacement", "R+ vs R-")
    ep, em = g("e", "whisker x displacement", "R+"), g("e", "whisker x displacement", "R-")
    s = new(prs, "Geometry: the R- whisker response slides toward the miss level", "State space per cohort (learners): passive pre -> post, whisker (yellow) and auditory (blue); active hits / misses for reference")
    picture(s, c158["e"], 0.4, 1.3, w=8.2, h=4.8, scale=20)
    eqbox(s, 8.9, 1.3, 4.1, [("axes", E["ss"])])
    bullets(s, 8.9, 3.4, 4.1, 2.4, [f"Whisker shift along x: R+ {ep.mean_a:+.2f} (p = {H.pnum(ep.p_nonparam)}), R- {em.mean_a:+.2f} (p = {H.pnum(em.p_nonparam)}); "
                                    f"R+ vs R- p = {H.pnum(ex_.p_nonparam)} (Mann-Whitney).",
                                    "Both cohorts shrink along y (shared, exposure); the x shift has no session-time null: illustration."], size=12)
    ew, el = R["expo_dyW"], R["expo_lick_WR"]
    s = new(prs, "Exposure shrinks the responses but does not move them on the hit / miss axis", "R- mice receive more whisker and auditory trials, in the same proportion")
    picture(s, figs["expo"], 0.3, 1.3, w=9.0)
    eqbox(s, 9.5, 1.3, 3.5, [("covariate model", E["expo"])])
    bullets(s, 9.5, 2.45, 3.5, 3.0, [f"Whisker shrinkage: trial counts p = {H.pnum(ew[2])}, {H.pnum(ew[3])}; cohort p = {H.pnum(ew[1])}.",
                                     f"Lick-axis change: cohort p = {H.pnum(el[1])} with counts."], size=12)

    # ------------------------------------------------------------------ C
    D.section_slide(prs, "C", "Per area group", "Each area group with its own axes; >= 3 sessions per cohort, >= 20 tracked units")
    for f, t in ((PUB / "151_area_heatmap_learners.png", "Pre -> post changes per area group, learners (151)"),
                 (PUB / "151_area_heatmap_all.png", "Pre -> post changes per area group, all mice (151)"),
                 (PUB / "147_choice_axis_readout_area_groups.png", "Decoder readout change per area group (147)")):
        if f.exists():
            s = new(prs, t, "Columns: displacement along the hit / miss axis (whisker, auditory, W - A) and the excess over the linear-shift null; "
                            "colour scaled per column; uncorrected")
            picture(s, f, 0.3, 1.3, w=12.7, h=5.5, scale=20, center=True)
    s = new(prs, "Per area: the R- direction in most groups, strongest in whisker sensorimotor areas, midbrain and thalamus")
    bullets(s, 0.6, 1.5, 12.2, 5, ["Cohort differences with both tests in somatosensory-whisker cortex, motor areas, midbrain and thalamus (all mice).",
                                   "Most area groups point in the R- direction; per-area samples are small and tests uncorrected: descriptive.",
                                   "No area-by-cohort interaction test yet; the whole-brain result is the inference."], size=16)

    # ------------------------------------------------------------------ summary
    D.CUR["sec"] = "S"
    s = new(prs, "COSYNE figure: within-day change on the hit / miss axis (learners)")
    picture(s, PUB / "158_cosyne_within_day_learners.png", 0.3, 1.3, w=12.7, scale=20)
    bullets(s, 0.6, 4.4, 12.2, 2.6, ["a: before the task the whisker response lies between hits and misses; b, c: projection on the hit / miss axis per epoch, "
                                     "whisker and auditory; d: pre -> post change beyond session time; e: state space per cohort.",
                                     f"Whisker post - pre R+ vs R- p = {H.pnum(bpp.p_nonparam)}; beyond session time p = {H.pnum(dp.p_nonparam)} (whisker), "
                                     f"{H.pnum(dwap.p_nonparam)} (whisker - auditory); auditory n.s."], size=13)
    s = new(prs, "Summary")
    bullets(s, 0.6, 1.4, 12.2, 5.4, ["Choice is decodable from the first 50 ms in both cohorts, and decodability barely changes within the session.",
                                     "After the task, the R- whisker response loses its hit-side component on the active hit / miss axis; R+ does not.",
                                     "The change exceeds what session time produces (linear-shift null) and is whisker-specific (auditory control).",
                                     "Exposure explains the shared shrinkage, not the hit / miss shift.",
                                     "Area groups point the same way; strongest in whisker sensorimotor areas, midbrain and thalamus."], size=17)
    s = new(prs, "Caveats")
    bullets(s, 0.6, 1.4, 12.2, 5.4, ["Session time: inference rests on the excess over a conservative linear-shift null.",
                                     "State: the pre-stimulus window shifts miss-ward after the task in R- for both stimuli; the whisker-specific claim rests on "
                                     "whisker - auditory contrasts.",
                                     "Reward rate and contingency cannot be separated (R- receive fewer rewards by design).",
                                     "5-35 ms hit / miss axes are unreliable; raw cosines and projections are used.",
                                     "One session per mouse; exploratory and uncorrected; expert sessions not shown here."], size=16)
    D.CUR["sec"] = "bk"
    for f, t in ((PUB / "149_passive_readout_shift_null_all.png", "Decoder readout vs the shift null (149)"),
                 (PUB / "150_axis_alignment_shift_null_all.png", "Lick axis and late coding direction vs the shift null (150)"),
                 (PUB / "151_state_space_variants_learners.png", "State space, three y axes (151)"),
                 (PUB / "151_state_space_heatmap_learners.png", "State-space displacements (151)"),
                 (PUB / "153_exposure_control.png", "Exposure control (153)"), (PUB / "147_controls.png", "State and engagement controls (147)"),
                 (PUB / "138_cosyne_lt_placebo_all.png", "Learning trial vs placebo (138)"), (PUB / "144_hitmedian_split_all.png", "Hit-median vs midpoint split (144)")):
        if f.exists():
            s = new(prs, t); picture(s, f, 0.4, 1.2, w=12.5, h=5.8, scale=20, center=True)
    prs.save(OUT)
    print(OUT, D.CUR["n"], "slides")


if __name__ == "__main__":
    main()
