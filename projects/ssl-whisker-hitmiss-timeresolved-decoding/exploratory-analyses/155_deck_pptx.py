"""155 -- Main-results deck as PowerPoint (user 2026-10-05: "build a unique slide deck that describes the main results; equations
when needed; high quality figures"; "PowerPoint"; scope: this project + the within-day pre-lick result; Part A with midpoint,
hit-median and learning-trial splits and their controls; two slides of negative results).
Flow (one message per slide; methods introduced one at a time):
  intro      title; task and question
  A          A1 choice is decodable early (method: decoding vs linear-shift null); A2 midpoint vs hit-median split; A3 learning trial
             vs placebo splits; A4 robustness of the learning-trial effect (windows, step vs gradual); A5 within-day pre-lick convergence
  B          B1 one pattern, one axis (equations); B2 lick-axis alignment across epochs; B3 session-time confound (linear-shift null);
             B4 beyond session time with the auditory control; B5 robustness (learners, late coding direction); B6 brain areas;
             B7 state space; B8 exposure
  C          summary across methods (forest plot); synthesis; caveats and next steps
  backup     full figures (138, 143, 144, 147 controls, 149, 150, 151 variants / heatmap, 153)
Figures: every panel re-plotted from the per-session tables with the publication helpers (143 / 149 / 153 / 154), exported at 600
dpi and placed at ~2.2x, so text is ~11-13 pt on the slide. Equations: matplotlib mathtext rendered to transparent PNGs.
Numbers on slides: computed here from the tables (and numbers.json of the report for session-wide decoding).
Output: combined_results_ks4/<slug>/report/main_results.pptx (+ figures/deck_pptx/*.png) and report.md, the slim main-results report
(same figures and numbers; the full report is report_full.md from report/build_report.py, run that first: numbers.json)
Run (haas, repo root): python .../155_deck_pptx.py
"""

from __future__ import annotations

import importlib
import json
import re
import sys
import textwrap
import warnings
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Rectangle
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

warnings.filterwarnings("ignore")
EA = Path(__file__).resolve().parent
sys.path.insert(0, str(EA)); sys.path.insert(0, str(EA.parents[2] / "scripts"))
H = importlib.import_module("143_lt_split_windows_figures")
F149 = importlib.import_module("149_passive_readout_shift_null_figure")
F154 = importlib.import_module("154_cosyne_remapping_figure")
from axel_bisi_paths import axel_bisi_root  # noqa: E402

COL, COH = H.COL, H.COH
WC, AC = "#f7b519", "#2c2cdb"
FIG, PUB = EA / "figures", EA / "figures" / "publication"
OUTF = FIG / "deck_pptx"
RDIR = axel_bisi_root() / "combined_results_ks4" / "ssl-whisker-hitmiss-timeresolved-decoding" / "report"
PPTX = RDIR / "main_results.pptx"
DPI, SCALE = 600, 2.2
# 2026-10-05 user: "plain style, just black text on white, boxes for take-home and that is it"
INK = SOFT = RGBColor(0, 0, 0)
SEC = {"intro": RGBColor(0x5C, 0x67, 0x7D), "A": RGBColor(0x2A, 0x9D, 0x8F), "B": RGBColor(0xE0, 0x7A, 0x3F),
       "C": RGBColor(0x7B, 0x5E, 0xA7), "bk": RGBColor(0x8D, 0x99, 0xAE)}
SEC_TAG = {"intro": "", "A": "A. Active trials: does choice decoding change?", "B": "B. Passive vs active: one method, made robust",
           "C": "C. Across methods", "bk": "Backup"}
CUR = {"sec": "intro", "n": 0}


def P(a, b):
    return f"{H.pnum(a)} | {H.pnum(b)}"


def perm_diff(x, y, n=20000, seed=0):
    rng = np.random.default_rng(seed); z = np.r_[x, y]; k = len(x); d0 = np.mean(x) - np.mean(y)
    null = np.array([np.mean(p[:k]) - np.mean(p[k:]) for p in (rng.permutation(z) for _ in range(n))])
    return (np.sum(np.abs(null) >= abs(d0)) + 1) / (n + 1)


# ------------------------------------------------------------------------------------------------------------- figure helpers
def savefig(fig, name):
    OUTF.mkdir(parents=True, exist_ok=True)
    p = OUTF / f"{name}.png"
    fig.savefig(p, dpi=DPI, bbox_inches="tight", pad_inches=0.03, facecolor="white")
    plt.close(fig)
    return p


def eq(name, lines, fs=13):
    """mathtext equation lines -> transparent PNG (placed at scale 1)"""
    fig = plt.figure(figsize=(6, 0.5 * len(lines)))
    for i, t in enumerate(lines):
        fig.text(0.0, 1 - (i + 0.5) / len(lines), t, fontsize=fs, va="center", color="black")
    OUTF.mkdir(parents=True, exist_ok=True)
    p = OUTF / f"eq_{name}.png"
    fig.savefig(p, dpi=DPI, bbox_inches="tight", pad_inches=0.04, transparent=True)
    plt.close(fig)
    return p


def pair_panel(ax, a, b, labels, ylab, title, R, key):
    """before / after per cohort (open -> filled), paired tests per cohort, cohort test on the change"""
    xp = {"R+": (0, 1), "R-": (2.4, 3.4)}
    H.mean_pair(ax, xp, a, b)
    out = {}
    for c in COH:
        pw, pt, n = H.paired(b[c], a[c])
        out[c] = (np.nanmean(b[c] - a[c]), pw, pt, n)
        ax.text(np.mean(xp[c]), 1.02, P(pw, pt), color=COL[c], ha="center", fontsize=5, transform=ax.get_xaxis_transform())
    mw, we = H.unpaired(b["R+"] - a["R+"], b["R-"] - a["R-"])
    out["cohort"] = (mw, we)
    ax.text(1.7, 1.12, f"R+ vs R- {P(mw, we)}", ha="center", fontsize=5, transform=ax.get_xaxis_transform())
    ax.set_xticks([0, 1, 2.4, 3.4]); ax.set_xticklabels(labels * 2, fontsize=5.2); ax.set_xlim(-0.5, 3.9)
    for t, c in zip(ax.get_xticklabels(), ["R+", "R+", "R-", "R-"]):
        t.set_color(COL[c])
    ax.set_ylabel(ylab); ax.set_title(title, pad=24)
    R[key] = out


def dots_panel(ax, v, ylab, title, R, key, perm=False):
    for k, c in enumerate(COH):
        ax.plot(k + np.random.default_rng(k).uniform(-0.14, 0.14, len(v[c])), v[c], "o", ms=2, color=COL[c], alpha=0.45, mew=0)
        ax.errorbar(k, np.nanmean(v[c]), H.sem(v[c]), fmt="o", ms=3.8, color=COL[c], lw=1.1, capsize=0, zorder=5)
        pw, pt, n = H.one_sample(v[c])
        ax.text(k, 1.02, P(pw, pt), color=COL[c], ha="center", fontsize=5, transform=ax.get_xaxis_transform())
        R[f"{key}_{c}"] = (np.nanmean(v[c]), pw, pt, n)
    mw, we = H.unpaired(v["R+"], v["R-"]); pm = perm_diff(v["R+"], v["R-"]) if perm else np.nan
    R[f"{key}_cohort"] = (mw, we, pm)
    ax.text(0.5, 1.12, f"R+ vs R- {P(mw, we)}" + (f"; perm. {H.pnum(pm)}" if perm else ""), ha="center", fontsize=5,
            transform=ax.get_xaxis_transform())
    ax.axhline(0, color="0.6", lw=0.5, ls=(0, (2, 2)))
    ax.set_xticks([0, 1]); ax.set_xticklabels(["R+", "R-"]); ax.set_xlim(-0.6, 1.6)
    for t, c in zip(ax.get_xticklabels(), COH):
        t.set_color(COL[c])
    ax.set_ylabel(ylab); ax.set_title(title, pad=24)


# ------------------------------------------------------------------------------------------------------------- figures
def fig_task():
    fig, ax = plt.subplots(figsize=(3.2, 2.2)); ax.axis("off"); ax.set_xlim(0, 1); ax.set_ylim(0, 1)
    for i, (c, t) in enumerate((("R+", "whisker  ->  lick  ->  reward"), ("R-", "whisker  ->  lick  ->  no reward"))):
        ax.text(0.08, 0.75 - i * 0.3, c, color=COL[c], fontsize=10, weight="bold", va="center")
        ax.text(0.25, 0.75 - i * 0.3, t, fontsize=7, va="center")
    ax.text(0.25, 0.15, "auditory  ->  lick  ->  reward  (both cohorts)", fontsize=7, va="center")
    return savefig(fig, "task")


def fig_timeline():
    fig, ax = plt.subplots(figsize=(4.2, 1.1)); ax.axis("off"); ax.set_xlim(0, 1); ax.set_ylim(0, 1)
    for x0, x1, t, c in ((0, 0.18, "passive pre", "0.88"), (0.2, 0.8, "active: whisker, auditory, catch trials", "#d6ece6"),
                         (0.82, 1, "passive post", "0.88")):
        ax.plot([x0 + 0.005, x1 - 0.005], [0.45, 0.45], color="black", lw=2.5, solid_capstyle="butt")
        ax.text((x0 + x1) / 2, 0.62, t, ha="center", va="center", fontsize=6.5)
    ax.annotate("", xy=(1, 0.25), xytext=(0, 0.25), arrowprops=dict(arrowstyle="->", color="black", lw=0.8))
    ax.text(0.5, 0.05, "learning session (whisker day 0), one per mouse", ha="center", fontsize=5.8, color="black")
    return savefig(fig, "timeline")


def fig_decode_schematic():
    fig, ax = plt.subplots(figsize=(3.0, 1.6)); ax.axis("off"); ax.set_xlim(0, 1); ax.set_ylim(0, 1)
    ax.text(0.18, 0.71, "active whisker trials\nhit | miss", ha="center", va="center", fontsize=6.5)
    ax.annotate("", xy=(0.55, 0.71), xytext=(0.38, 0.71), arrowprops=dict(arrowstyle="->", lw=1.2))
    ax.text(0.77, 0.71, "cross-validated\nlinear decoder", ha="center", va="center", fontsize=6.5)
    ax.annotate("", xy=(0.77, 0.32), xytext=(0.77, 0.53), arrowprops=dict(arrowstyle="->", lw=1.2))
    ax.text(0.77, 0.2, "accuracy minus\nlinear-shift null", ha="center", va="center", fontsize=6.5, weight="bold")
    ax.text(0.0, 0.2, "labels shifted 10-50 %\nagainst the trials (no wrap)", fontsize=5.5, color="black", va="center")
    return savefig(fig, "decode_schematic")


def fig_splits(R):
    D = pd.read_parquet(EA / "139_hitmedian_split_whole_brain.parquet")
    D = D[D.skipped_reason.isna() & (D.decoding == "hitmiss") & (D.window == "5-100ms")]
    fig, axes = plt.subplots(1, 3, figsize=(5.4, 1.9), gridspec_kw=dict(width_ratios=[1, 1, 1], wspace=0.75))
    for j, (sp, t) in enumerate((("mid", "midpoint split"), ("hitmedian", "hit-median split"))):
        d = D[D.split == sp]
        a = {c: d[d.reward_group == c].sep_corr_1.to_numpy(float) for c in COH}
        b = {c: d[d.reward_group == c].sep_corr_2.to_numpy(float) for c in COH}
        pair_panel(axes[j], a, b, ["before", "after"], "hit / miss decoding\n(acc - shift null)" if j == 0 else "", t, R, f"split_{sp}")
    d = D[D.split == "hitmedian"]; dm = D[D.split == "mid"]
    a = {c: dm[dm.reward_group == c].hit_rate_1.to_numpy(float) for c in COH}
    b = {c: dm[dm.reward_group == c].hit_rate_2.to_numpy(float) for c in COH}
    pair_panel(axes[2], a, b, ["1st", "2nd"], "whisker hit rate", "behaviour (midpoint halves)", R, "split_hr")
    return savefig(fig, "splits")


def fig_lt(R):
    A, S, _ = H.load("5-100ms", "_A1v2", "_step1_pl100", "all")
    a_ = A[A.definition == "L0 stored"]; s_ = S[S.variant == "L0 stored"]
    fig, axes = plt.subplots(1, 3, figsize=(5.4, 1.9), gridspec_kw=dict(wspace=0.75))
    a = {c: a_[a_.reward_group == c].corr_pre_matched.to_numpy(float) for c in COH}
    b = {c: a_[a_.reward_group == c].corr_post_matched.to_numpy(float) for c in COH}
    pair_panel(axes[0], a, b, ["before", "after"], "hit / miss decoding\n(acc - shift null)", "split at the learning trial", R, "lt_ba")
    a = {c: s_[s_.reward_group == c].placebo_mean.to_numpy(float) for c in COH}
    b = {c: s_[s_.reward_group == c].delta.to_numpy(float) for c in COH}
    pair_panel(axes[1], a, b, ["placebo", "real"], "after - before", "real vs placebo splits", R, "lt_rp")
    v = {c: (s_[s_.reward_group == c].delta - s_[s_.reward_group == c].placebo_mean).to_numpy(float) for c in COH}
    dots_panel(axes[2], v, "change beyond placebo", "excess", R, "lt_ex", perm=True)
    return savefig(fig, "lt")


def fig_lt_robust(R):
    fig, axes = plt.subplots(1, 2, figsize=(4.6, 1.9), gridspec_kw=dict(width_ratios=[1.5, 1], wspace=0.55))
    ax = axes[0]
    for i, (win, t118, t122, wl) in enumerate(H.WINS):
        _, S, _ = H.load(win, t118, t122, "all")
        s_ = S[S.variant == "L5 whisker CP"]
        v = {c: (s_[s_.reward_group == c].delta - s_[s_.reward_group == c].placebo_mean).to_numpy(float) for c in COH}
        for k, c in enumerate(COH):
            x = i + (k - 0.5) * 0.3
            ax.errorbar(x, np.nanmean(v[c]), H.sem(v[c]), fmt="o", ms=3.8, color=COL[c], lw=1.1, capsize=0)
            pw, pt, n = H.one_sample(v[c]); R[f"l5_{win}_{c}"] = (np.nanmean(v[c]), pw, pt, n)
        mw, we = H.unpaired(v["R+"], v["R-"]); R[f"l5_{win}_cohort"] = (mw, we)
        ax.text(i, 1.02, P(mw, we), ha="center", fontsize=4.8, transform=ax.get_xaxis_transform())
    ax.axhline(0, color="0.6", lw=0.5, ls=(0, (2, 2)))
    ax.set_xticks(range(len(H.WINS))); ax.set_xticklabels([w[0].replace("ms", " ms") for w in H.WINS]); ax.set_xlim(-0.5, len(H.WINS) - 0.5)
    ax.set_ylabel("change beyond placebo"); ax.set_title("L5 change point, three windows\n(R+ vs R-)", pad=12)
    # step vs gradual R+ (5-50 ms)
    ax = axes[1]
    win, t118, t122, _ = [w for w in H.WINS if w[0] == "5-50ms"][0]
    A, S, _ = H.load(win, t118, t122, "all")
    L = H.learners(); step_sids = set(S[S.variant == "L5 whisker CP"].session_id) | set(A[A.definition == "L5 whisker CP"].session_id)
    g = {"step": S[S.variant == "L5 whisker CP"], "gradual": S[(S.variant == "half") & ~S.session_id.isin(step_sids)]}
    g["gradual"] = g["gradual"][H.mouse_of(g["gradual"]).isin(L)]
    vv = {}
    for k, (gn, d) in enumerate(g.items()):
        v = (d[d.reward_group == "R+"].delta - d[d.reward_group == "R+"].placebo_mean).to_numpy(float); vv[gn] = v
        ax.plot(k + np.random.default_rng(k).uniform(-0.14, 0.14, len(v)), v, "o", ms=2, color=COL["R+"], alpha=0.45, mew=0)
        ax.errorbar(k, np.nanmean(v), H.sem(v), fmt="o", ms=3.8, color=COL["R+"], mfc=COL["R+"] if gn == "step" else "white", lw=1.1, capsize=0)
    mw, we = H.unpaired(vv["step"], vv["gradual"]); R["sg"] = (mw, we, len(vv["step"]), len(vv["gradual"]))
    ax.text(0.5, 1.02, P(mw, we), ha="center", fontsize=4.8, transform=ax.get_xaxis_transform())
    ax.axhline(0, color="0.6", lw=0.5, ls=(0, (2, 2))); ax.set_xticks([0, 1]); ax.set_xticklabels(["step", "gradual"]); ax.set_xlim(-0.6, 1.6)
    ax.set_title("R+ learners, 5-50 ms", pad=12)
    return savefig(fig, "lt_robust")


def fig_wd(R):
    TR = pd.read_csv(F154.WD / F154.WD_SLOPES / "trial_slopes_trajectories.csv"); D2 = pd.read_csv(F154.WD / F154.WD_SLOPES / "trial_slopes_sessions.csv")
    fig, axes = plt.subplots(1, 2, figsize=(4.4, 1.9), gridspec_kw=dict(width_ratios=[1.8, 1], wspace=0.6))
    rows = []
    F154.panel_b(axes[0], TR, rows); F154.panel_c(axes[1], D2, rows)
    axes[0].set_title(axes[0].get_title(), fontsize=axes[0].title.get_fontsize(), pad=10)
    R["wd"] = pd.DataFrame(rows)
    return savefig(fig, "wd")


def fig_pattern_axis():
    fig, ax = plt.subplots(figsize=(2.2, 1.8)); ax.axis("off"); ax.set_xlim(-0.15, 1.25); ax.set_ylim(-0.25, 1.05); ax.set_aspect("equal")
    ax.add_patch(FancyArrowPatch((0, 0), (1.1, 0), arrowstyle="-|>", mutation_scale=9, color="0.15", lw=1.4))
    ax.text(1.05, -0.12, "lick axis L", ha="center", fontsize=6.5)
    p = (0.6, 0.72); ax.add_patch(FancyArrowPatch((0, 0), p, arrowstyle="-|>", mutation_scale=9, color=WC, lw=1.6))
    ax.text(0.62, 0.84, "passive whisker\npattern p", color="#c48a00", ha="center", fontsize=6)
    ax.plot([p[0], p[0]], [0, p[1]], color="0.5", lw=0.6, ls=":"); ax.plot([0, p[0]], [-0.04, -0.04], color="0.3", lw=2.5)
    ax.text(0.3, -0.2, r"projection $\pi$", ha="center", fontsize=6, color="0.3"); ax.text(0.15, 0.06, r"$\theta$", fontsize=8)
    return savefig(fig, "pattern_axis")


def fig_null():
    fig, ax = plt.subplots(figsize=(4.0, 1.6)); ax.axis("off"); ax.set_xlim(-2, 100); ax.set_ylim(-0.3, 4.2)
    rng = np.random.default_rng(3); n = 50; y = rng.random(n) < np.linspace(0.8, 0.15, n); xs = np.linspace(0, 62, n); k = 15
    ax.scatter(xs, np.full(n, 3.4), c=["0.15" if h else "0.78" for h in y], s=9, marker="s")
    ax.text(64, 3.4, "real labels (hit dark, miss light)", va="center", fontsize=5.8)
    ax.scatter(xs, np.full(n, 2.5), c="#9ecae1", s=9, marker="s"); ax.text(64, 2.5, "neural trials", va="center", fontsize=5.8)
    ax.scatter(xs[: n - k], np.full(n - k, 1.3), c=["0.15" if h else "0.78" for h in y[k:]], s=9, marker="s")
    ax.scatter(xs[: n - k], np.full(n - k, 0.5), c="#9ecae1", s=9, marker="s")
    ax.text(64, 0.9, "null: labels shifted by k trials,\nno wrap-around (hits still early)", va="center", fontsize=5.8)
    ax.annotate("", xy=(xs[0], 1.65), xytext=(xs[k], 3.1), arrowprops=dict(arrowstyle="->", color=AC, lw=0.9))
    ax.text(xs[k // 2] - 3, 2.05, "shift k", color=AC, fontsize=5.5)
    return savefig(fig, "null")


def load_135():
    A = pd.read_parquet(EA / "135_alignment_epochs_tracked.parquet")
    return A[(A.area == "All units") & A.skipped_reason.isna()]


def fig_lick_epochs(R):
    fig, ax = plt.subplots(figsize=(2.9, 2.0))
    rows = []; F154.panel_d(ax, load_135(), rows); R["d"] = pd.DataFrame(rows)
    ax.set_title("whisker (solid) and auditory (dashed)\npattern vs the active lick axis")
    return savefig(fig, "lick_epochs")


def excess_fig(name, specs, R, w=5.2):
    fig, axes = plt.subplots(1, len(specs), figsize=(w, 2.0), gridspec_kw=dict(wspace=0.75))
    rows = []
    for ax, (d, col, title) in zip(np.atleast_1d(axes), specs):
        F149.excess_panel(ax, d, col, title, rows, "", name)
        ax.set_title(ax.get_title(), fontsize=ax.title.get_fontsize(), pad=26)
    np.atleast_1d(axes)[0].set_ylabel("pre -> post change\nbeyond session time")
    R[name] = pd.DataFrame(rows)
    return savefig(fig, name)


def fig_state(R):
    D = pd.read_parquet(EA / "146_choice_axis_readout.parquet")
    D = D[D.skipped_reason.isna() & (D.area == "whole_brain") & (D.unit_set == "stable") & (D.response == "epochbase")]
    fig, ax = plt.subplots(figsize=(2.7, 2.1))
    rows = []; F154.panel_f(ax, D, rows); R["f"] = pd.DataFrame(rows)
    for t in list(ax.texts):
        if t.get_text().startswith("dashed"):
            t.remove()  # stated in the slide subtitle
    return savefig(fig, "state")


def fig_exposure(R):
    F153 = importlib.import_module("153_exposure_control")
    X = F153.load()
    fig, axes = plt.subplots(1, 3, figsize=(5.6, 1.9), gridspec_kw=dict(wspace=0.6))
    rows = []
    F153.corr_panel(axes[0], X, "dyW", "whisker response, own axis\n(shrinkage)", rows, "s")
    F153.corr_panel(axes[1], X, "lick_WR", "whisker vs lick axis\n(excess)", rows, "c")
    axes[0].set_ylabel("post - pre")
    import statsmodels.formula.api as smf
    ax = axes[2]
    meas = [("dyW", "W shrinkage"), ("dyA", "A shrinkage"), ("lick_WR", "lick axis W"), ("lick_WAP", "lick axis W-A"), ("dx_W", "state x, W")]
    for i, (y, lab) in enumerate(meas):
        g = X.dropna(subset=[y]); sd = g[y].std()
        for k, (form, mk) in enumerate(((f"{y} ~ Rm", "o"), (f"{y} ~ Rm + nWz + nAz", "s"))):
            f = smf.ols(form, g).fit(); ci = f.conf_int().loc["Rm"]
            ax.errorbar(f.params.Rm / sd, i + (0.15 if k else -0.15), xerr=[[(f.params.Rm - ci[0]) / sd], [(ci[1] - f.params.Rm) / sd]], fmt=mk,
                        ms=3, color="k" if k else "0.6", mfc="white" if k else "0.6", lw=0.8, capsize=0)
            if k:
                R[f"expo_{y}"] = (f.params.Rm / sd, f.pvalues.Rm, f.pvalues.nWz, f.pvalues.nAz)
    ax.axvline(0, color="0.6", lw=0.5, ls=(0, (2, 2))); ax.set_yticks(range(len(meas))); ax.set_yticklabels([m[1] for m in meas], fontsize=5.2)
    ax.invert_yaxis(); ax.set_xlabel("R- minus R+ (SD)"); ax.set_title("cohort effect: without (grey) /\nwith trial counts (black)")
    R["expo_rows"] = pd.DataFrame(rows)
    return savefig(fig, "exposure")


def fig_forest(R, A, C, V, D, WD2):
    """standardised R- minus R+ difference (pooled SD) with bootstrap 95 % CI, across methods"""
    rng = np.random.default_rng(0)
    items = [("pre-lick drift toward rewarded licks (within-day)", WD2[WD2.stage == "learning"], "md_WH-FA_slope", "cohort"),
             ("lick axis: whisker cosine", A, "shift_excess_dWR", "reward_group"), ("lick axis: whisker - auditory projection", A, "shift_excess_dWAP", "reward_group"),
             ("late coding direction: passive axis cosine", C, "shift_excess_daxisR", "reward_group"),
             ("late coding direction: whisker - auditory projection", C, "shift_excess_dWAP", "reward_group"),
             ("decoder: whisker - auditory (per-trial baseline)", V, "shift_excess_dWA", "reward_group"),
             ("state space: whisker along the choice axis", D, "dx_W", "reward_group")]
    fig, ax = plt.subplots(figsize=(4.2, 2.3))
    for i, (lab, d, col, gc) in enumerate(items):
        x, y = d[d[gc] == "R+"][col].dropna().to_numpy(float), d[d[gc] == "R-"][col].dropna().to_numpy(float)
        sd = np.sqrt(((len(x) - 1) * x.var(ddof=1) + (len(y) - 1) * y.var(ddof=1)) / (len(x) + len(y) - 2))
        eff = (y.mean() - x.mean()) / sd
        bs = [(rng.choice(y, len(y)).mean() - rng.choice(x, len(x)).mean()) / sd for _ in range(3000)]
        lo, hi = np.percentile(bs, [2.5, 97.5])
        mw, we = H.unpaired(x, y)
        colr = "#e07a3f" if "within-day" not in lab else "#2a9d8f"
        ax.errorbar(eff, i, xerr=[[eff - lo], [hi - eff]], fmt="o", ms=4, color=colr, lw=1.2, capsize=0)
        ax.text(1.02, i, P(mw, we), transform=ax.get_yaxis_transform(), va="center", fontsize=5)
        R[f"forest_{i}"] = (lab, eff, lo, hi, mw, we, len(x), len(y))
    ax.axvline(0, color="0.5", lw=0.6, ls=(0, (2, 2)))
    ax.set_yticks(range(len(items))); ax.set_yticklabels([t[0] for t in items], fontsize=5.6); ax.invert_yaxis()
    ax.set_xlabel("R- minus R+ (pooled SD; 95 % bootstrap CI)\nnegative = away from the lick / reward-lick representation")
    ax.text(1.02, -0.9, "MW | Welch", transform=ax.get_yaxis_transform(), fontsize=5, weight="bold")
    return savefig(fig, "forest")


def fig_synthesis():
    fig, axes = plt.subplots(1, 2, figsize=(4.8, 2.0))
    for ax, (title, lab_axis, a0, a_rp, a_rm, labs) in zip(axes, (
            ("pre-lick (100 ms before the lick)", "SL -> AH (reward-lick)", 40, 25, 75, ("WH day-0 start", "R+ late", "R- late")),
            ("stimulus onset (5-35 ms, passive)", "lick axis (hit - miss)", 35, 30, 80, ("whisker pre", "R+ post", "R- post")))):
        ax.axis("off"); ax.set_xlim(-0.2, 1.3); ax.set_ylim(-0.3, 1.1); ax.set_aspect("equal")
        ax.add_patch(FancyArrowPatch((0, 0), (1.15, 0), arrowstyle="-|>", mutation_scale=8, color="0.15", lw=1.3))
        ax.text(0.95, -0.15, lab_axis, ha="center", fontsize=5.5)
        for ang, colr, ls, lab in ((a0, WC, "-", labs[0]), (a_rp, COL["R+"], (0, (3, 2)), labs[1]), (a_rm, COL["R-"], (0, (3, 2)), labs[2])):
            r = np.deg2rad(ang); e = (0.85 * np.cos(r), 0.85 * np.sin(r))
            ax.add_patch(FancyArrowPatch((0, 0), e, arrowstyle="-|>", mutation_scale=8, color=colr, lw=1.3, linestyle=ls))
            ax.text(e[0] * 1.18, e[1] * 1.12, lab, color=colr, fontsize=5.3, ha="center")
        ax.set_title(title, fontsize=6.5)
    return savefig(fig, "synthesis")


# ------------------------------------------------------------------------------------------------------------- pptx helpers
def pfix(t):
    """'p = <0.001' -> 'p < 0.001'"""
    return re.sub(r"= <\s?", "< ", t)


def rect(slide, x, y, w, h, color, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = color
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, s, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    paras = s if isinstance(s, list) else [s]
    for i, t in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = pfix(t); r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold; r.font.name = "Calibri"
    return tb


def bullets(slide, x, y, w, h, items, size=13):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    for i, t in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        r = p.add_run(); r.text = "•  " + pfix(t); r.font.size = Pt(size); r.font.color.rgb = INK; r.font.name = "Calibri"
    return tb


def picture(slide, path, x, y, w=None, h=None, scale=SCALE, center=False):
    """place a 600-dpi PNG at `scale` times its native size (or fit into w / h); center: centre horizontally in w"""
    im = Image.open(path); wi, hi = im.size[0] / DPI * scale, im.size[1] / DPI * scale
    if w is not None and wi > w:
        hi, wi = hi * w / wi, w
    if h is not None and hi > h:
        wi, hi = wi * h / hi, h
    if center and w is not None:
        x = x + (w - wi) / 2
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(wi), Inches(hi))
    return wi, hi


def takehome(slide, x, y, w, h, s):
    b = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF); b.line.color.rgb = INK; b.line.width = Pt(1); b.shadow.inherit = False
    tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(0.18); tf.margin_top = Inches(0.12)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT; r = p.add_run(); r.text = "Take-home"; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = INK; r.font.name = "Calibri"
    p = tf.add_paragraph(); p.alignment = PP_ALIGN.LEFT; p.space_before = Pt(4); r = p.add_run(); r.text = pfix(s); r.font.size = Pt(13); r.font.color.rgb = INK; r.font.name = "Calibri"


def eqbox(slide, x, y, w, items, title="DEFINITIONS"):
    """items: list of (label, eq png); returns the height used"""
    hs = []
    for lab, png in items:
        im = Image.open(png); hs.append(im.size[1] / DPI)
    lh = [0.07 + 0.2 * int(np.ceil(len(lab) / (12.5 * (w - 0.2)))) for lab, _ in items]  # ~12.5 characters per inch at 10 pt
    h = sum(l + hh for l, hh in zip(lh, hs)) + 0.05
    yy = y
    for (lab, png), hh, l in zip(items, hs, lh):
        text(slide, x + 0.12, yy, w - 0.2, l, lab, size=10, color=SOFT)
        picture(slide, png, x + 0.15, yy + l - 0.02, w=w - 0.3, scale=1.0)
        yy += l + hh
    return h


def new_slide(prs, title, sub=None):
    CUR["n"] += 1
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if SEC_TAG[CUR["sec"]]:
        text(s, 0.35, 0.05, 10, 0.3, SEC_TAG[CUR["sec"]], size=10)
    text(s, 0.35, 0.28, 12.6, 0.5, title, size=24, bold=True)
    if sub:
        text(s, 0.35, 0.7, 12.6, 0.35, sub, size=13, color=SOFT)
    text(s, 12.4, 7.12, 0.6, 0.3, str(CUR["n"]), size=9, align=PP_ALIGN.RIGHT)
    return s


def section_slide(prs, key, title, sub):
    CUR["sec"] = key; CUR["n"] += 1
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.9, 2.6, 11.5, 1.2, f"{key}.  {title}", size=32, bold=True)
    text(s, 0.9, 3.9, 11.5, 1.2, sub, size=18)


# ------------------------------------------------------------------------------------------------------------- main
def main():
    H.setup()
    plt.rcParams.update({"mathtext.fontset": "dejavusans"})
    R = {}
    N = json.loads((RDIR / "numbers.json").read_text())
    A = load_135()
    C = pd.read_parquet(EA / "140_coding_direction_noise.parquet"); C = C[(C.split == "hitmedian") & C.skipped_reason.isna()]
    V = pd.read_parquet(EA / "146_choice_axis_readout.parquet")
    V = V[(V.area == "whole_brain") & (V.unit_set == "stable") & V.skipped_reason.isna()]
    VD = V[V.response == "epochbase"].copy(); VD["dx_W"] = VD.ss_post_W_x - VD.ss_pre_W_x
    WD2 = pd.read_csv(F154.WD / F154.WD_SLOPES / "trial_slopes_sessions.csv")
    L = H.learners(); AL = A[H.mouse_of(A).isin(L)]
    figs = dict(task=fig_task(), timeline=fig_timeline(), dec=fig_decode_schematic(), splits=fig_splits(R), lt=fig_lt(R), ltr=fig_lt_robust(R),
                wd=fig_wd(R), pa=fig_pattern_axis(), null=fig_null(), epochs=fig_lick_epochs(R),
                b4=excess_fig("b4", [(A, "shift_excess_dWR", "whisker\ncosine"), (A, "shift_excess_dAR", "auditory\ncosine"),
                                     (A, "shift_excess_dWAP", "whisker - auditory\nprojection")], R),
                b5=excess_fig("b5", [(AL, "shift_excess_dWR", "learners:\nwhisker cosine"), (AL, "shift_excess_dWAP", "learners:\nW - A projection"),
                                     (C, "shift_excess_daxisR", "late CD: passive\naxis cosine"), (C, "shift_excess_dWAP", "late CD:\nW - A projection")], R, w=7.6),
                state=fig_state(R), expo=fig_exposure(R), forest=fig_forest(R, A, C, V[V.response == "trialbase"], VD, WD2), syn=fig_synthesis())
    E = dict(
        dec=eq("dec", [r"$\Delta\mathrm{acc} = \mathrm{acc}_{\mathrm{real}} - \langle \mathrm{acc}^{(k)}_{\mathrm{shift}} \rangle_k$"]),
        split=eq("split", [r"midpoint: $t_{n/2}$;   hit-median: hit $\lfloor H/2 \rfloor + 1$"]),
        lt=eq("lt", [r"$\delta_s = \Delta\mathrm{acc}^{\mathrm{after}}_s - \Delta\mathrm{acc}^{\mathrm{before}}_s$",
                     r"$\mathrm{excess} = \delta_{\mathrm{LT}} - \langle \delta_s \rangle_{s \in \mathrm{placebo}}$"]),
        cd=eq("cd", [r"$c_i = \dfrac{\mathbf{x}_i\cdot\widehat{\mathbf{CD}} - \langle \cdot \rangle_{\mathrm{SL}}}{\langle \cdot \rangle_{\mathrm{AH}} - \langle \cdot \rangle_{\mathrm{SL}}}$,   $\widehat{\mathbf{CD}} \propto \bar{\mathbf{x}}_{\mathrm{AH}} - \bar{\mathbf{x}}_{\mathrm{SL}}$"]),
        drift=eq("drift", [r"$c_i = a_k + b_k\,\tau_i,\quad \beta = b_{\mathrm{WH}} - b_{\mathrm{SL}}$"]),
        resp=eq("resp", [r"$r_{it} = f^{[5,35]}_{it} - \langle f^{[-55,-20]}_{it'} \rangle_{t' \in \mathrm{epoch}},\quad \varepsilon_{it} = r_{it}/\sigma_i$"]),
        lick=eq("lick", [r"$\mathbf{L} = \bar{\mathbf{z}}_{\mathrm{hit}} - \bar{\mathbf{z}}_{\mathrm{miss}},\quad \mathbf{p}_W = \langle \boldsymbol{\varepsilon}_t \rangle_{t \in W}$"]),
        cos=eq("cos", [r"$\cos\theta = \dfrac{\mathbf{p}\cdot\mathbf{L}}{\|\mathbf{p}\|\,\|\mathbf{L}\|},\quad \pi = \dfrac{\mathbf{p}\cdot\hat{\mathbf{L}}}{\sqrt{n}}$"]),
        excess=eq("excess", [r"$\Delta M = M^{\mathrm{post}} - M^{\mathrm{pre}}$",
                             r"$\mathrm{excess} = \Delta M - \dfrac{1}{K}\sum_{k=1}^{K}\Delta M^{(k)}_{\mathrm{shift}}$"]),
        contrast=eq("contrast", [r"$\Delta\pi_W - \Delta\pi_A$  (drift shared by both stimuli cancels)"]),
        cd2=eq("cd2", [r"$\mathbf{CD}_2 = \bar{\mathbf{z}}_{\mathrm{hit},2} - \bar{\mathbf{z}}_{\mathrm{miss},2}$"]),
        ss=eq("ss", [r"$\hat{\mathbf{x}} = \mathbf{L}/\|\mathbf{L}\|,\quad \hat{\mathbf{y}} \propto \mathbf{p}^{\mathrm{pre}}_W - (\mathbf{p}^{\mathrm{pre}}_W\cdot\hat{\mathbf{x}})\,\hat{\mathbf{x}}$",
                     r"$\Delta x = (\mathbf{p}^{\mathrm{post}} - \mathbf{p}^{\mathrm{pre}})\cdot\hat{\mathbf{x}}$"]),
        expo=eq("expo", [r"$\Delta M = \beta_0 + \beta_{R-}\,\mathbb{1}_{R-} + \beta_W z(n_W) + \beta_A z(n_A) + \epsilon$"]))

    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    # ---- title
    CUR["sec"] = "intro"; CUR["n"] += 1
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.7, 0.8, 12, 1.6, "Within one learning session, the reward contingency remaps the whisker sensorimotor chain", size=32, bold=True)
    text(s, 0.7, 2.45, 12, 0.5, "Whole-brain Neuropixels during single-session whisker learning; R+ (whisker licks rewarded) vs R- (not rewarded)",
         size=16)
    for i, (k, t) in enumerate((("A", "Active trials: is choice decoding stable within the session? Where does it change?"),
                                ("B", "Passive vs active: does the early whisker response move relative to the lick axis? One method, made robust"),
                                ("C", "Across methods: summary, synthesis, caveats"))):
        text(s, 0.7, 3.9 + i * 1.0, 0.75, 0.6, k, size=24, bold=True)
        text(s, 1.4, 3.95 + i * 1.0, 11.3, 0.8, t, size=18)
    # ---- task
    s = new_slide(prs, "Same stimuli, same action, different contingency", "Learning session (whisker day 0), one per mouse")
    picture(s, figs["task"], 0.5, 1.4, w=6.2); picture(s, figs["timeline"], 0.5, 5.3, w=6.2)
    bullets(s, 7.2, 1.5, 5.7, 4.5, ["Hit = lick after the whisker stimulus: same label in both cohorts (trained response in R+, error in R-).",
                                    "H1: choice information grows when R+ mice learn; H2: not in R-.",
                                    "H3: the early whisker response is re-mapped relative to the lick representation, depending on the contingency.",
                                    "Statistics: session = unit (one per mouse); a non-parametric and a parametric test (shown as p | p); uncorrected."], size=15)
    # ---- A
    section_slide(prs, "A", "Active trials: is choice decoding stable within the session?", "One method throughout: hit / miss decoding against a linear-shift null")
    s = new_slide(prs, "Choice is decodable early after the whisker, equally in both cohorts", "Method A: decode hit vs miss from the 5-50 ms response; subtract a time-preserving null")
    picture(s, figs["dec"], 0.5, 1.4, w=6.4)
    eqbox(s, 0.5, 5.0, 6.4, [("plotted: accuracy above the linear-shift null (k = 10-50 % of the trials)", E["dec"])])
    k0 = lambda k: N.get(k, np.nan)
    bullets(s, 7.3, 1.5, 5.6, 3.5, [f"Session-wide, 5-50 ms: R+ {k0('k0_rp'):+.3f}, R- {k0('k0_rm'):+.3f} above the null.",
                                    f"No cohort difference (p = {H.pnum(k0('k0_mw'))} | {H.pnum(k0('k0_w'))}).",
                                    "Question: does this change within the session, and differently in R+ and R-?"], size=15)
    takehome(s, 7.3, 5.0, 5.6, 1.4, "A lick after the whisker is predictable from early whole-brain activity, rewarded or not.")
    sm, sh, hr = R["split_mid"], R["split_hitmedian"], R["split_hr"]
    s = new_slide(prs, "Splitting the session in two: no change in choice decoding", "Midpoint split (time) and hit-median split (equal hits before and after); hit / miss 5-100 ms")
    picture(s, figs["splits"], 0.3, 1.3, w=8.6)
    eqbox(s, 9.2, 1.3, 3.8, [("split points", E["split"])])
    bullets(s, 9.2, 2.6, 3.8, 3.0, [f"Midpoint: R+ {sm['R+'][0]:+.3f}, R- {sm['R-'][0]:+.3f}; R+ vs R- p = {P(*sm['cohort'])}.",
                                    f"Hit-median: R+ {sh['R+'][0]:+.3f}, R- {sh['R-'][0]:+.3f}; R+ vs R- p = {P(*sh['cohort'])}.",
                                    f"Control: hit rate changes between halves (R+ p = {P(hr['R+'][1], hr['R+'][2])}, R- p = {P(hr['R-'][1], hr['R-'][2])}); "
                                    "the hit-median split equalises the number of hits per half, not the hit rate."], size=13)
    takehome(s, 9.2, 5.6, 3.8, 1.3, "Decodability is stable between halves in both cohorts.")
    lb, lr, le = R["lt_ba"], R["lt_rp"], R
    s = new_slide(prs, "At the learning trial, R+ decoding rises beyond placebo splits", "Split at the behavioural learning trial; control: every other split of the same session")
    picture(s, figs["lt"], 0.3, 1.3, w=8.6)
    eqbox(s, 9.2, 1.3, 3.8, [("change at a split s and excess over placebo splits (>= 10 trials away)", E["lt"])])
    bullets(s, 9.2, 3.25, 3.8, 2.4, [f"R+ excess {le['lt_ex_R+'][0]:+.3f} (p = {P(le['lt_ex_R+'][1], le['lt_ex_R+'][2])}); R- {le['lt_ex_R-'][0]:+.3f}.",
                                     f"R+ vs R- p = {P(le['lt_ex_cohort'][0], le['lt_ex_cohort'][1])}; cohort-label permutation p = {H.pnum(le['lt_ex_cohort'][2])}."], size=13)
    takehome(s, 9.2, 5.6, 3.8, 1.3, "A small R+-specific gain appears when the split is tied to behaviour.")
    s = new_slide(prs, "The learning-trial gain holds across windows and belongs to abrupt learners", "L5: Bayesian change point of the whisker hit sequence")
    picture(s, figs["ltr"], 0.4, 1.3, w=8.4)
    bl = [f"{w[0].replace('ms', ' ms')}: R+ {R[f'l5_{w[0]}_R+'][0]:+.3f} (p = {P(R[f'l5_{w[0]}_R+'][1], R[f'l5_{w[0]}_R+'][2])}); R+ vs R- p = "
          f"{P(*R[f'l5_{w[0]}_cohort'])}" for w in H.WINS]
    bullets(s, 9.2, 1.4, 3.8, 3.5, bl + [f"R+ step ({R['sg'][2]}) vs gradual ({R['sg'][3]}) learners: p = {P(R['sg'][0], R['sg'][1])}."], size=12)
    takehome(s, 9.2, 5.4, 3.8, 1.5, "The change-point effect is real but small, and carried by mice that learn abruptly.")
    wd = R["wd"]; c_ = wd[(wd.panel == "c") & (wd.cohort == "R+ vs R-")].iloc[0]; cp_ = wd[(wd.panel == "c") & (wd.cohort == "R+")].iloc[0]
    cm_ = wd[(wd.panel == "c") & (wd.cohort == "R-")].iloc[0]
    s = new_slide(prs, "Pre-lick activity of whisker hits converges on rewarded licks in R+, diverges in R-", "100 ms before the first lick; position on the session's reward-lick coding direction (within-day project)")
    picture(s, figs["wd"], 0.3, 1.3, w=8.6)
    eqbox(s, 9.2, 1.3, 3.8, [("projection (SL = 0, AH = 1)", E["cd"]), ("per-session drift over normalised time", E["drift"])])
    bullets(s, 9.2, 3.95, 3.8, 1.6, [f"Drift R+ {cp_.mean_a:+.2f} (p = {P(cp_.p_nonparam, cp_.p_param)}), R- {cm_.mean_a:+.2f} (p = {P(cm_.p_nonparam, cm_.p_param)}).",
                                     f"R+ vs R- p = {P(c_.p_nonparam, c_.p_param)}; {c_.note}. Sessions with >= 4 whisker hits."], size=12)
    takehome(s, 9.2, 5.6, 3.8, 1.3, "Decodability is stable, but the pre-lick signal moves in opposite directions.")
    # ---- B
    section_slide(prs, "B", "Passive vs active: one method, made robust", "Does the earliest (5-35 ms) whisker response move relative to the active lick axis, before vs after the task?")
    s = new_slide(prs, "One pattern, one axis", "Passive whisker pattern p (before or after the task) against the active lick axis L")
    picture(s, figs["pa"], 0.6, 1.4, w=4.6); picture(s, figs["timeline"], 0.4, 5.4, w=5.6)
    eqbox(s, 6.4, 1.3, 6.6, [("response of unit i on trial t (5-35 ms minus the epoch's baseline), evoked scaling", E["resp"]),
                              ("lick axis (active whisker trials) and passive whisker pattern", E["lick"]), ("direction and reach", E["cos"])])
    bullets(s, 6.4, 5.1, 6.6, 1.8, ["Same tracked, drift-checked units in every epoch; active trials with a lick before 35 ms excluded.",
                                    "Prediction (H3): after the task, the R- whisker pattern points less along L; R+ unchanged; auditory unchanged."], size=13)
    d = R["d"].iloc[0]
    s = new_slide(prs, "After the task, the R- whisker response leaves the lick axis", "Cosine of the evoked patterns with the lick axis, in the four epochs")
    picture(s, figs["epochs"], 0.6, 1.3, w=7.6)
    bullets(s, 8.6, 1.5, 4.4, 3.0, [f"Passive post - pre, R+ vs R-: p = {P(d.p_nonparam, d.p_param)}.",
                                    "Auditory pattern (dashed): no cohort difference.",
                                    "Active halves are descriptive (they average hits and misses); inference uses passive pre -> post."], size=14)
    takehome(s, 8.6, 5.2, 4.4, 1.5, "R- whisker responses decouple from the lick axis during the task and stay decoupled.")
    s = new_slide(prs, "Session time is a confound, and a linear-shift null removes it", "Hit / miss labels drift (R- hits early, misses late); any hit / miss axis can partly encode early vs late")
    picture(s, figs["null"], 0.5, 1.5, w=7.8)
    eqbox(s, 8.7, 1.3, 4.3, [("change of a metric, and excess over K = 50 axes rebuilt from shifted labels", E["excess"])])
    bullets(s, 0.6, 5.0, 12, 1.9, ["Shuffled labels destroy trial order: a decoder that learned 'late = miss' is not in that null.",
                                   "A circular shift turns the trend into a sawtooth and centres the null near 0: too lenient.",
                                   "The linear shift keeps the drift of both series; conservative, since real learning is time-correlated too."], size=14)
    b4 = R["b4"]; g4 = lambda m: b4[(b4.measure == m) & (b4.cohort == "R+ vs R-")].iloc[0]
    s = new_slide(prs, "The R- decoupling survives session time and the auditory control", "Passive pre -> post change minus the change produced by lick axes rebuilt from shifted labels (135)")
    picture(s, figs["b4"], 0.3, 1.3, w=8.6)
    eqbox(s, 9.2, 1.3, 3.8, [("linear contrast", E["contrast"])])
    bullets(s, 9.2, 2.5, 3.8, 3.0, [f"Whisker cosine R+ vs R- p = {P(g4('shift_excess_dWR').p_nonparam, g4('shift_excess_dWR').p_param)}.",
                                    f"Auditory cosine p = {P(g4('shift_excess_dAR').p_nonparam, g4('shift_excess_dAR').p_param)}.",
                                    f"Whisker - auditory projection p = {P(g4('shift_excess_dWAP').p_nonparam, g4('shift_excess_dWAP').p_param)}."], size=13)
    takehome(s, 9.2, 5.4, 3.8, 1.5, "Beyond session time, only the R- whisker response moves away from the lick axis.")
    b5 = R["b5"]; g5 = lambda i: b5[b5.cohort == "R+ vs R-"].iloc[i]
    s = new_slide(prs, "Robust in learners only and with a second axis", "Learners only (lick axis); late hit / miss coding direction of a hit-median split (140)")
    picture(s, figs["b5"], 0.3, 1.3, w=9.4)
    eqbox(s, 9.9, 1.3, 3.1, [("late coding direction", E["cd2"])])
    bullets(s, 9.9, 2.4, 3.1, 3.2, [f"Learners: p = {P(g5(0).p_nonparam, g5(0).p_param)} and {P(g5(1).p_nonparam, g5(1).p_param)}.",
                                    f"Late CD: p = {P(g5(2).p_nonparam, g5(2).p_param)} and {P(g5(3).p_nonparam, g5(3).p_param)}."], size=12)
    takehome(s, 9.9, 5.4, 3.1, 1.5, "Same answer in learners and with a different hit / miss axis.")
    s = new_slide(prs, "Strongest in whisker sensorimotor areas, midbrain and thalamus", "Each area group with its own axes; >= 3 sessions per cohort, >= 20 units")
    picture(s, PUB / "151_area_heatmap_all.png", 0.3, 1.2, w=12.7, h=5.1, scale=20, center=True)
    text(s, 0.4, 6.35, 12.5, 0.7, "Cohort differences with both tests in somatosensory-whisker cortex, motor areas, midbrain and thalamus; "
         "most areas point in the R- direction; single cells are small and uncorrected.", size=13, color=SOFT)
    f_ = R["f"].iloc[0]
    s = new_slide(prs, "Geometry: the R- whisker response slides to the miss level", "State space, passive pre at the origin; dashed: where active misses sit on the choice axis")
    picture(s, figs["state"], 0.8, 1.3, w=6.8)
    eqbox(s, 8.0, 1.3, 5.0, [("axes", E["ss"])])
    bullets(s, 8.0, 3.45, 5.0, 2.0, [f"Displacement along the choice axis, R+ vs R-: p = {P(f_.p_nonparam, f_.p_param)}.",
                                     "Both cohorts shrink along the whisker pattern (shared)."], size=13)
    takehome(s, 8.0, 5.3, 5.0, 1.4, "After the task, an R- whisker stimulus evokes a pattern like a whisker trial without a lick.")
    ew, el = R["expo_dyW"], R["expo_lick_WR"]
    s = new_slide(prs, "Exposure shrinks the responses but does not move them on the choice axis", "R- mice receive more whisker and auditory trials, in the same proportion")
    picture(s, figs["expo"], 0.3, 1.3, w=9.0)
    eqbox(s, 9.5, 1.3, 3.5, [("covariate model", E["expo"])])
    bullets(s, 9.5, 2.45, 3.5, 3.0, [f"Whisker shrinkage: trial counts p = {H.pnum(ew[2])} (whisker), {H.pnum(ew[3])} (auditory); cohort p = {H.pnum(ew[1])}.",
                                     f"Lick-axis change: cohort p = {H.pnum(el[1])} with counts; counts p = {H.pnum(el[2])}, {H.pnum(el[3])}."], size=12)
    takehome(s, 9.5, 5.3, 3.5, 1.5, "Adaptation-like shrinkage and contingency-specific re-mapping are separable.")
    # ---- C
    section_slide(prs, "C", "Across methods", "Summary, synthesis, caveats")
    s = new_slide(prs, "The same direction across methods and windows", "Standardised R- minus R+ difference; teal: pre-lick (within-day), orange: stimulus onset beyond session time")
    picture(s, figs["forest"], 0.6, 1.3, w=9.5)
    bullets(s, 10.3, 1.5, 2.7, 4.5, ["Mean-difference axes, the decoder relative to auditory, and the state space agree.",
                                     "The decoder's whisker readout alone does not separate the cohorts beyond session time (backup)."], size=13)
    s = new_slide(prs, "Synthesis", "Decodability is stable; the representation relative to the choice is re-mapped, in opposite directions")
    picture(s, figs["syn"], 0.6, 1.4, w=8.4)
    bullets(s, 9.2, 1.5, 3.8, 4.5, ["R+: pre-lick whisker hits converge on the reward-lick representation within the learning session.",
                                    "R-: pre-lick whisker hits diverge from it, and the earliest whisker response decouples from the lick axis.",
                                    "Neither is visible in decoding accuracy, which stays stable."], size=14)
    takehome(s, 0.6, 5.7, 12.2, 1.1, "The reward contingency remaps the whisker sensorimotor chain within one session; decodability alone misses it.")
    s = new_slide(prs, "Caveats and next steps")
    bullets(s, 0.6, 1.4, 12.2, 5.4, ["Session time: hit / miss labels drift; part B inference is on the excess over a conservative linear-shift null.",
                                     "State: the pre-stimulus window shifts miss-ward after the task in R- (both stimuli); the whisker-specific claim rests on "
                                     "whisker - auditory contrasts.",
                                     "Exposure and reward: R- receive more trials (same proportion) and fewer rewards (by design); exposure scales shrinkage, not the "
                                     "choice-axis change; reward rate and contingency cannot be separated.",
                                     "Noise: 5-35 ms hit / miss axes are unreliable; tests use raw cosines and projections.",
                                     "Units differ between the pre-lick (good + mua) and stimulus-onset (tracked stable) analyses.",
                                     "Exploratory, uncorrected. Next: held-out mice or expert sessions; time-matched decoder; area-resolved state space."], size=16)
    # ---- backup
    CUR["sec"] = "bk"
    for path, t in ((PUB / "138_cosyne_lt_placebo_all.png", "Learning trial vs placebo, full figure (138)"),
                    (PUB / "143_lt_split_L5_windows_all.png", "L5 change point across windows (143)"),
                    (PUB / "143_step_vs_gradual_all.png", "Step vs gradual learners (143)"),
                    (PUB / "144_hitmedian_split_all.png", "Hit-median vs midpoint split, all windows (144)"),
                    (PUB / "149_passive_readout_shift_null_all.png", "Decoder readout vs the shift null; state readout (149)"),
                    (PUB / "150_axis_alignment_shift_null_all.png", "Lick axis and late coding direction vs the shift null (150)"),
                    (PUB / "151_state_space_variants_all.png", "State space, three y axes (151)"),
                    (PUB / "151_state_space_heatmap_all.png", "State-space displacements (151)"),
                    (PUB / "153_exposure_control.png", "Exposure control, full figure (153)"),
                    (PUB / "147_controls.png", "State and engagement controls (147)")):
        if path.exists():
            s = new_slide(prs, t)
            picture(s, path, 0.4, 1.2, w=12.5, h=5.8, scale=20, center=True)
    PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(PPTX)
    print(PPTX, CUR["n"], "slides")
    write_report(R, N, figs)


# ------------------------------------------------------------------------------------------------------------- slim report
def write_report(R, N, figs):
    """report.md: the main results only, same figures and numbers as the deck (the full report is report_full.md, build_report.py)"""
    import shutil
    from datetime import date
    (RDIR / "figures").mkdir(parents=True, exist_ok=True)
    for k, p in figs.items():
        shutil.copyfile(p, RDIR / "figures" / f"main_{p.name}")
    shutil.copyfile(PUB / "151_area_heatmap_all.png", RDIR / "figures" / "main_151_area_heatmap_all.png")
    F = lambda k, cap, w=100: f"![{cap}](figures/main_{figs[k].name}){{width={w}%}}\n"
    pp = lambda a, b: f"$p$ = {H.pnum(a)} | {H.pnum(b)}"
    k0 = lambda k: N.get(k, np.nan)
    sm, sh, hr = R["split_mid"], R["split_hitmedian"], R["split_hr"]
    wd = R["wd"]; wc = {c: wd[(wd.panel == "c") & (wd.cohort == c)].iloc[0] for c in ("R+", "R-", "R+ vs R-")}
    wb = {c: wd[(wd.panel == "b") & (wd.cohort == c)].iloc[0] for c in COH}
    d_ = R["d"].iloc[0]; f_ = R["f"].iloc[0]
    b4, b5 = R["b4"], R["b5"]
    g4 = lambda m, c="R+ vs R-": b4[(b4.measure == m) & (b4.cohort == c)].iloc[0]
    g5 = lambda i: b5[b5.cohort == "R+ vs R-"].iloc[i]
    ew, ea, el, elw = R["expo_dyW"], R["expo_dyA"], R["expo_lick_WR"], R["expo_lick_WAP"]
    er = R["expo_rows"]; erw = {c: er[(er.measure == "dyW") & (er.cohort == c) & (er.x == "nW")].iloc[0] for c in COH}
    wins = [w[0] for w in H.WINS]
    fo = [R[f"forest_{i}"] for i in range(7)]
    md = f"""---
title: "Within one learning session, the reward contingency remaps the whisker sensorimotor chain"
subtitle: "Main results (ssl-whisker-hitmiss-timeresolved-decoding + within-day pre-lick); full report: report_full.pdf"
date: "{date.today():%Y-%m-%d}"
geometry: margin=2.2cm
fontsize: 10pt
colorlinks: true
header-includes:
  - \\usepackage{{caption}}
  - \\captionsetup{{font=small}}
  - \\usepackage{{float}}
  - \\floatplacement{{figure}}{{H}}
---

**Data.** SSL dataset (private), Kilosort 4 (`ssl_ephys`), whisker day 0 (learning session), one session per mouse; R+ mice are rewarded
for licking after the whisker stimulus, R- mice are not. Trials with `perf == 6` excluded (passive trials kept). **Statistics.** Session
= unit; every comparison reports a non-parametric and a parametric test as $p$ = non-parametric | parametric (one sample: Wilcoxon |
$t$; paired: Wilcoxon | paired $t$; R+ vs R-: Mann-Whitney | Welch); uncorrected, exploratory. Slides: `main_results.pptx` (same figures
and numbers).

**Summary.** (A) Hit / miss (lick vs no lick after the whisker) is decodable from the 5-50 ms whole-brain response, equally in R+ and
R-, and this decodability does not change between session halves. It rises in R+ only when the session is split at the behavioural
learning trial, a small effect carried by abrupt learners. Pre-lick activity of whisker hits converges on the reward-lick representation
in R+ and diverges from it in R-. (B) Comparing passive whisker responses before and after the task, the R- whisker response decouples
from the active lick axis beyond what session time explains, the auditory response does not, and the result holds in learners, with a
second hit / miss axis, in a state-space view and with trial counts as covariates. Decodability is stable; the representation relative
to the choice is re-mapped, in opposite directions in the two cohorts.

# A. Active trials: is choice decoding stable within the session?

## A1. Choice is decodable early, in both cohorts

Hit vs miss is decoded from the 5-50 ms response (cross-validated linear decoder). Hit / miss labels drift over the session (R- hits are
early, misses late), so the reference is a **linear-shift null**: labels are shifted by $k$ trials against the neural trials, 10-50 % of
the trials, without wrap-around, which keeps the slow trend of both series and destroys only the trial-by-trial match:

$$\\Delta\\mathrm{{acc}} = \\mathrm{{acc}}_{{\\mathrm{{real}}}} - \\big\\langle \\mathrm{{acc}}^{{(k)}}_{{\\mathrm{{shift}}}} \\big\\rangle_k .$$

Session-wide, R+ decode {k0('k0_rp'):+.3f} and R- {k0('k0_rm'):+.3f} above the null; no cohort difference ({pp(k0('k0_mw'), k0('k0_w'))}).

## A2. Splitting the session in two: no change

{F('splits', 'Session halves. Hit / miss decoding (5-100 ms, accuracy minus shift null) before vs after the split, per cohort (open: before, filled: after); right: whisker hit rate in the two midpoint halves. Above each panel: R+ vs R- test on the change; coloured: paired test within cohort.')}
Midpoint split (time; $t_{{n/2}}$): change R+ {sm['R+'][0]:+.3f}, R- {sm['R-'][0]:+.3f}, R+ vs R- {pp(*sm['cohort'])}. Hit-median split
(after hit $\\lfloor H/2 \\rfloor + 1$, equal hits per half): R+ {sh['R+'][0]:+.3f}, R- {sh['R-'][0]:+.3f}, R+ vs R- {pp(*sh['cohort'])}. The hit
rate falls between halves in both cohorts (R+ {pp(hr['R+'][1], hr['R+'][2])}, R- {pp(hr['R-'][1], hr['R-'][2])}); the hit-median split
equalises the number of hits per half, not the hit rate.

## A3. At the learning trial, R+ decoding rises beyond placebo splits

{F('lt', 'Split at the learning trial (L0). Left: before vs after; middle: change at the real split vs the mean change at placebo splits (every other split of the session, >= 10 trials from the real one); right: excess per session.')}
With $\\delta_s = \\Delta\\mathrm{{acc}}^{{\\mathrm{{after}}}}_s - \\Delta\\mathrm{{acc}}^{{\\mathrm{{before}}}}_s$ the change at a split $s$,

$$\\mathrm{{excess}} = \\delta_{{\\mathrm{{LT}}}} - \\langle \\delta_s \\rangle_{{s \\in \\mathrm{{placebo}}}} .$$

R+ excess {R['lt_ex_R+'][0]:+.3f} ({pp(R['lt_ex_R+'][1], R['lt_ex_R+'][2])}), R- {R['lt_ex_R-'][0]:+.3f}; R+ vs R-
{pp(R['lt_ex_cohort'][0], R['lt_ex_cohort'][1])}, cohort-label permutation $p$ = {H.pnum(R['lt_ex_cohort'][2])}.

## A4. The learning-trial gain holds across windows and belongs to abrupt learners

{F('ltr', 'Left: excess over placebo at the L5 change point (Bayesian change point of the whisker hit sequence), three windows; numbers: R+ vs R-. Right: R+ learners with a change point (step) vs learners without (gradual, half split), 5-50 ms.', 85)}
""" + "".join(f"- {w.replace('ms', ' ms')}: R+ {R[f'l5_{w}_R+'][0]:+.3f} ({pp(R[f'l5_{w}_R+'][1], R[f'l5_{w}_R+'][2])}); R+ vs R- "
              f"{pp(*R[f'l5_{w}_cohort'])}.\n" for w in wins) + f"""- R+ step ({R['sg'][2]}) vs gradual ({R['sg'][3]}) learners: {pp(R['sg'][0], R['sg'][1])}.

The change-point effect is real but small, and it is carried by mice that learn abruptly.

## A5. Pre-lick activity of whisker hits converges on rewarded licks in R+, diverges in R-

{F('wd', 'Within-day project (100 ms before the first lick). Left: whisker hits (WH) minus spontaneous licks (SL) on the reward-lick coding direction, five bins of learning-session time (curves: all day-0 sessions), and all expert sessions (dots, unpaired). Right: per-session drift of WH toward auditory hits (AH). Sessions with >= 4 whisker hits.', 85)}
Each session's reward-lick coding direction $\\widehat{{\\mathbf{{CD}}}} \\propto \\bar{{\\mathbf{{x}}}}_{{\\mathrm{{AH}}}} - \\bar{{\\mathbf{{x}}}}_{{\\mathrm{{SL}}}}$;
trial $i$ is placed at

$$c_i = \\frac{{\\mathbf{{x}}_i\\cdot\\widehat{{\\mathbf{{CD}}}} - \\langle\\cdot\\rangle_{{\\mathrm{{SL}}}}}}{{\\langle\\cdot\\rangle_{{\\mathrm{{AH}}}} - \\langle\\cdot\\rangle_{{\\mathrm{{SL}}}}}}, \\qquad
c_i = a_k + b_k\\,\\tau_i, \\qquad \\beta = b_{{\\mathrm{{WH}}}} - b_{{\\mathrm{{SL}}}},$$

with $\\tau$ normalised session time and $\\beta$ the drift of whisker hits relative to spontaneous licks. Drift R+
{wc['R+'].mean_a:+.2f} ({pp(wc['R+'].p_nonparam, wc['R+'].p_param)}), R- {wc['R-'].mean_a:+.2f} ({pp(wc['R-'].p_nonparam, wc['R-'].p_param)});
R+ vs R- {pp(wc['R+ vs R-'].p_nonparam, wc['R+ vs R-'].p_param)} ({wc['R+ vs R-'].note}). Expert sessions: R+ {wb['R+'].note}, R- {wb['R-'].note}.

# B. Passive vs active: does the earliest whisker response move relative to the lick axis?

## B1. One pattern, one axis

Same tracked, drift-checked stable units in passive pre, both active halves and passive post (137b). Response of unit $i$ on trial
$t$: 5-35 ms rate minus the epoch's pre-stimulus baseline (-55 to -20 ms), scaled by the unit's evoked SD,
$\\varepsilon_{{it}} = r_{{it}}/\\sigma_i$. The **lick axis** is the active hit - miss difference of whisker trials and the **pattern** the
mean passive response,

$$\\mathbf{{L}} = \\bar{{\\mathbf{{z}}}}_{{\\mathrm{{hit}}}} - \\bar{{\\mathbf{{z}}}}_{{\\mathrm{{miss}}}}, \\qquad \\mathbf{{p}}_W = \\langle \\boldsymbol{{\\varepsilon}}_t \\rangle_{{t \\in W}}, \\qquad
\\cos\\theta = \\frac{{\\mathbf{{p}}\\cdot\\mathbf{{L}}}}{{\\|\\mathbf{{p}}\\|\\,\\|\\mathbf{{L}}\\|}}, \\qquad \\pi = \\frac{{\\mathbf{{p}}\\cdot\\hat{{\\mathbf{{L}}}}}}{{\\sqrt{{n}}}} .$$

Prediction: after the task the R- whisker pattern points less along $\\mathbf{{L}}$; R+ unchanged; the auditory pattern unchanged.

## B2. After the task, the R- whisker response leaves the lick axis

{F('epochs', 'Cosine of the whisker (solid) and auditory (dashed) patterns with the active lick axis in the four epochs (mean +- SEM over sessions).', 55)}
Passive post - pre, R+ vs R-: {pp(d_.p_nonparam, d_.p_param)}. Active halves are descriptive (their patterns average hits and misses);
inference uses passive pre -> post.

## B3. Beyond session time, and relative to the auditory response

Any hit / miss axis can partly encode early vs late. Each metric's change is therefore compared with the change produced by $K = 50$
lick axes rebuilt from labels shifted against the trials (same linear shift as A1):

$$\\Delta M = M^{{\\mathrm{{post}}}} - M^{{\\mathrm{{pre}}}}, \\qquad \\mathrm{{excess}} = \\Delta M - \\frac{{1}}{{K}}\\sum_{{k=1}}^{{K}} \\Delta M^{{(k)}}_{{\\mathrm{{shift}}}} .$$

The whisker - auditory contrast $\\Delta\\pi_W - \\Delta\\pi_A$ cancels any change shared by both stimuli (state, drift).

{F('b4', 'Excess over the linear-shift null, passive pre -> post (dots: sessions). Coloured: excess vs 0 per cohort; black: R+ vs R-.', 75)}
Whisker cosine: R+ {g4('shift_excess_dWR', 'R+').mean_a:+.3f}, R- {g4('shift_excess_dWR', 'R-').mean_a:+.3f}, R+ vs R-
{pp(g4('shift_excess_dWR').p_nonparam, g4('shift_excess_dWR').p_param)}. Auditory cosine: {pp(g4('shift_excess_dAR').p_nonparam, g4('shift_excess_dAR').p_param)}.
Whisker - auditory projection: {pp(g4('shift_excess_dWAP').p_nonparam, g4('shift_excess_dWAP').p_param)}.

## B4. Robust in learners and with a second hit / miss axis

{F('b5', 'Same excess, learners only (lick axis), and for the late coding direction $\\mathbf{CD}_2 = \\bar{\\mathbf{z}}_{\\mathrm{hit},2} - \\bar{\\mathbf{z}}_{\\mathrm{miss},2}$ of a hit-median split (140).', 90)}
Learners: whisker cosine {pp(g5(0).p_nonparam, g5(0).p_param)}, whisker - auditory projection {pp(g5(1).p_nonparam, g5(1).p_param)}. Late coding
direction: passive axis cosine {pp(g5(2).p_nonparam, g5(2).p_param)}, whisker - auditory projection {pp(g5(3).p_nonparam, g5(3).p_param)}.

## B5. Brain areas

![Each area group with its own axes (>= 3 sessions per cohort, >= 20 units). Left / middle: post - pre per cohort; right: R- minus R+; stars: one or both tests $p$ < 0.05; colour scaled per column.](figures/main_151_area_heatmap_all.png){{width=100%}}

The R- direction appears in most area groups; cohort differences with both tests in somatosensory-whisker cortex, motor areas, midbrain
and thalamus (single cells small, uncorrected).

## B6. Geometry: the R- whisker response slides to the miss level

{F('state', 'State space with passive pre at the origin: x = choice (lick) axis, y = passive-pre whisker pattern orthogonalised to x. Arrows: mean post - pre displacement per cohort (+ SEM); dots: sessions; dashed: active whisker misses on the choice axis.', 55)}
$$\\hat{{\\mathbf{{x}}}} = \\mathbf{{L}}/\\|\\mathbf{{L}}\\|, \\qquad \\hat{{\\mathbf{{y}}}} \\propto \\mathbf{{p}}^{{\\mathrm{{pre}}}}_W - (\\mathbf{{p}}^{{\\mathrm{{pre}}}}_W\\cdot\\hat{{\\mathbf{{x}}}})\\,\\hat{{\\mathbf{{x}}}}, \\qquad
\\Delta x = (\\mathbf{{p}}^{{\\mathrm{{post}}}} - \\mathbf{{p}}^{{\\mathrm{{pre}}}})\\cdot\\hat{{\\mathbf{{x}}}} .$$

Displacement along the choice axis: R+ {f_.mean_a:+.2f}, R- {f_.mean_b:+.2f}, {pp(f_.p_nonparam, f_.p_param)}. Both cohorts shrink along the
whisker pattern (shared). After the task, an R- whisker stimulus evokes a pattern like a whisker trial without a lick.

## B7. Exposure shrinks the responses but does not move them on the choice axis

R- sessions are longer: R- mice receive more whisker and auditory trials, in the same proportion. Per session,

$$\\Delta M = \\beta_0 + \\beta_{{R-}}\\,\\mathbf{{1}}_{{R-}} + \\beta_W\\, z(n_W) + \\beta_A\\, z(n_A) + \\epsilon .$$

{F('expo', 'Left: shrinkage of the whisker response along its own pattern vs active whisker trials; middle: lick-axis excess vs whisker trials (OLS line + 95 % CI, solid if p < 0.05); right: R- coefficient (SD units, 95 % CI) without (grey) and with (black) trial counts.')}
Whisker shrinkage scales with exposure (R- Spearman $\\rho$ = {erw['R-'].spearman:+.2f}, {pp(erw['R-'].p_nonparam, erw['R-'].p_param)}; model: whisker
trials $p$ = {H.pnum(ew[2])}, auditory trials $p$ = {H.pnum(ew[3])}, cohort $p$ = {H.pnum(ew[1])}). The lick-axis changes do not: with both counts
as covariates the cohort effect remains (whisker cosine $p$ = {H.pnum(el[1])}, whisker - auditory projection $p$ = {H.pnum(elw[1])}), counts
$p$ = {H.pnum(el[2])}, {H.pnum(el[3])}. Adaptation-like shrinkage and contingency-specific re-mapping are separable.

# C. Across methods

{F('forest', 'Standardised R- minus R+ difference (pooled SD, 95 % bootstrap CI); teal: pre-lick (within-day), orange: stimulus onset beyond session time; right: Mann-Whitney | Welch.', 85)}
""" + "".join(f"- {lab}: {eff:+.2f} SD [{lo:+.2f}, {hi:+.2f}], {pp(mw, we)} (R+ {nx}, R- {ny}).\n" for lab, eff, lo, hi, mw, we, nx, ny in fo) + f"""
The decoder's whisker readout alone does not separate the cohorts beyond session time (full report, III.5); its whisker - auditory
contrast does.

{F('syn', 'Synthesis (schematic).', 75)}
**Synthesis.** R+: pre-lick whisker hits converge on the reward-lick representation within the learning session. R-: pre-lick whisker
hits diverge from it, and the earliest whisker response decouples from the lick axis. Neither is visible in decoding accuracy, which
stays stable.

# Caveats

- Session time: hit / miss labels drift; part B inference is on the excess over a conservative linear-shift null (real learning is
  time-correlated too).
- State: the pre-stimulus window shifts miss-ward after the task in R- (both stimuli); the whisker-specific claim rests on the
  whisker - auditory contrasts.
- Exposure and reward: R- receive more trials (same proportion) and fewer rewards (by design); exposure scales shrinkage, not the
  choice-axis change; reward rate and contingency cannot be separated.
- Noise: 5-35 ms hit / miss axes are unreliable (noise-corrected cosines are a sensitivity check only); tests use raw cosines and
  projections.
- Units differ between the pre-lick (good + mua) and stimulus-onset (tracked stable) analyses.
- Exploratory and uncorrected. Next: held-out mice or expert sessions; time-matched decoder; area-resolved state space.

Everything else (session-wide time courses, pseudo-population, all split variants and windows, noise-corrected cosines, decoder readout,
controls 147, state-space variants) is in `report_full.pdf`.
"""
    md = re.sub(r"= <\s?", "< ", md)
    (RDIR / "report.md").write_text(md, encoding="utf-8")
    print(RDIR / "report.md")


if __name__ == "__main__":
    main()
