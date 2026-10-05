"""157 -- Story deck (user 2026-10-06: "make a separate slide deck with all necessary figures for this story").
Story: choice decoding is stable within the learning session, but the reward contingency re-maps the whisker sensorimotor chain at
both ends. Output end (pre-lick, within-day project): R+ whisker hits converge on rewarded licks, R- diverge. Input end (5-35 ms
stimulus onset, passive pre -> post): the R- whisker response decouples from the lick axis beyond session time, relative to the
auditory response; R+ does not change. The two analyses use the same logic (position of whisker-driven activity on a behaviourally
defined axis, with a within-session reference that shares the drift: spontaneous licks / the auditory response).
Claims kept out of the main line (caveats only): the shrinkage along the response's own pattern (y axis not cross-validated) and the
state-space x shift as an independent test (no session-time null).
Plain style and helpers from 155 (figures re-plotted from the tables at 600 dpi; numbers computed here).
Output: combined_results_ks4/<slug>/report/remapping_story.pptx. Run (haas, repo root): python .../157_story_deck.py
"""

from __future__ import annotations

import importlib
import json
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.patches import FancyArrowPatch
from pptx import Presentation
from pptx.util import Inches

EA = Path(__file__).resolve().parent
sys.path.insert(0, str(EA)); sys.path.insert(0, str(EA.parents[2] / "scripts"))
D = importlib.import_module("155_deck_pptx")
X = importlib.import_module("deck_full")
H, COL, COH, P = D.H, D.COL, D.COH, D.P
OUT = D.RDIR / "remapping_story.pptx"
D.SEC_TAG.update({"1": "1. Output end: just before the lick", "2": "2. Input end: the earliest whisker response", "3": "3. Both ends"})


# ----------------------------------------------------------------------------------------------------------- figures
def fig_cd_schematic():
    """the reward-lick axis: spontaneous licks (0) -> auditory hits (1); whisker hits drift along it over the session"""
    fig, ax = plt.subplots(figsize=(3.6, 1.6)); ax.axis("off"); ax.set_xlim(-0.35, 1.35); ax.set_ylim(-0.55, 0.75)
    ax.add_patch(FancyArrowPatch((-0.25, 0), (1.3, 0), arrowstyle="-|>", mutation_scale=8, color="black", lw=1.2))
    for x, lab in ((0, "spontaneous lick (SL)\nnever rewarded: 0"), (1, "auditory hit (AH)\nalways rewarded: 1")):
        ax.plot(x, 0, "o", color="black", ms=4); ax.text(x, -0.13, lab, ha="center", va="top", fontsize=5.5)
    for c, y, x0, x1 in (("R+", 0.42, 0.25, 0.55), ("R-", 0.22, 0.2, -0.12)):
        ax.add_patch(FancyArrowPatch((x0, y), (x1, y), arrowstyle="-|>", mutation_scale=7, color=COL[c], lw=1.3))
        ax.text(x0 + (0.04 if x1 > x0 else -0.04), y + 0.09, f"{c} whisker hits over the session", color=COL[c], fontsize=5.3,
                ha="left" if x1 > x0 else "right")
    ax.text(0.5, 0.66, "100 ms before the first lick: every trial is a lick", ha="center", fontsize=5.8)
    return D.savefig(fig, "story_cd_schematic")


def fig_exposure_story(R):
    """exposure vs the lick-axis change only (the sensory-axis shrinkage is left out: its y axis is not cross-validated)"""
    import statsmodels.formula.api as smf
    F153 = importlib.import_module("153_exposure_control")
    Xd = F153.load()
    fig, axes = plt.subplots(1, 3, figsize=(5.6, 1.9), gridspec_kw=dict(width_ratios=[0.9, 1, 1.1], wspace=0.7))
    ax = axes[0]
    for j, c_ in enumerate(("nW", "nA")):
        for k, coh in enumerate(COH):
            v = Xd[Xd.reward_group == coh][c_].to_numpy(float); x0 = j * 2.4 + k
            ax.plot(x0 + np.random.default_rng(k).uniform(-0.15, 0.15, len(v)), v, "o", ms=2, color=COL[coh], alpha=0.45, mew=0)
            ax.errorbar(x0, np.mean(v), H.sem(v), fmt="o", ms=3.6, color=COL[coh], lw=1, capsize=0)
        mw, we = H.unpaired(Xd[Xd.reward_group == "R+"][c_], Xd[Xd.reward_group == "R-"][c_]); R[f"expo_n_{c_}"] = (mw, we)
        ax.text(j * 2.4 + 0.5, 1.02, P(mw, we), ha="center", fontsize=5, transform=ax.get_xaxis_transform())
    fr = Xd.groupby("reward_group").fracW.mean(); R["expo_frac"] = (fr["R+"], fr["R-"], *H.unpaired(Xd[Xd.Rm == 0].fracW, Xd[Xd.Rm == 1].fracW))
    ax.set_xticks([0.5, 2.9]); ax.set_xticklabels(["whisker", "auditory"]); ax.set_ylabel("active trials"); ax.set_title("trials per session", pad=14)
    rows = []
    F153.corr_panel(axes[1], Xd, "lick_WR", "lick axis: whisker cosine\n(excess)", rows, "c")
    R["expo_corr"] = pd.DataFrame(rows)
    ax = axes[2]; out = []
    meas = [("lick_WR", "lick axis:\nwhisker cosine"), ("lick_WAP", "lick axis:\nW - A projection")]
    for i, (y, lab) in enumerate(meas):
        g = Xd.dropna(subset=[y]); sd = g[y].std()
        for k, (form, mk) in enumerate(((f"{y} ~ Rm", "o"), (f"{y} ~ Rm + nWz + nAz", "s"))):
            f = smf.ols(form, g).fit(); ci = f.conf_int().loc["Rm"]
            ax.errorbar(f.params.Rm / sd, i + (0.15 if k else -0.15), xerr=[[(f.params.Rm - ci[0]) / sd], [(ci[1] - f.params.Rm) / sd]], fmt=mk,
                        ms=3, color="k" if k else "0.6", mfc="white" if k else "0.6", lw=0.8, capsize=0)
            out.append(dict(measure=y, model="+ counts" if k else "cohort only", coef=f.params.Rm / sd, p=f.pvalues.Rm,
                            p_nW=f.pvalues.get("nWz", np.nan), p_nA=f.pvalues.get("nAz", np.nan)))
    R["expo_ols"] = pd.DataFrame(out)
    ax.axvline(0, color="0.6", lw=0.5, ls=(0, (2, 2))); ax.set_yticks(range(len(meas))); ax.set_yticklabels([m[1] for m in meas], fontsize=5.2)
    ax.invert_yaxis(); ax.set_xlabel("R- minus R+ (SD, 95 % CI)"); ax.set_title("cohort effect without (grey) /\nwith trial counts (black)")
    return D.savefig(fig, "story_exposure")


def fig_expert_story(R):
    """learning vs expert for the four measures that carry the story (156 panels; session = unit)"""
    F156 = importlib.import_module("156_expert_control")
    data = {st: F156.load(st) for st in F156.STAGES}
    meas = [("135", "shift_excess_dWR", "lick axis: whisker cosine\n(excess)", "change"),
            ("135", "shift_excess_dWAP", "lick axis: W - A projection\n(excess)", "change"),
            ("146t", "shift_excess_dWA", "decoder: W - A readout\n(excess)", "change"),
            ("135", "evokedW_cos_passive_pre", "lick axis: whisker cosine\nat passive pre (level)", "level")]
    fig, axes = plt.subplots(1, 4, figsize=(7.4, 2.9), gridspec_kw=dict(wspace=0.55))
    rows, rng = [], np.random.default_rng(0)
    for ax, (src, col, title, kind) in zip(axes, meas):
        v = {(c, st): data[st][src][data[st][src].reward_group == c][col].to_numpy(float) for c in COH for st in F156.STAGES}
        F156.panel(ax, v, title, rows, kind, col, rng)
        ax.set_title(ax.get_title(), fontsize=5.6, pad=40)   # clear the two lines of test labels above the axes
    R["expert"] = pd.DataFrame(rows)
    return D.savefig(fig, "story_expert")


def fig_chain_table(R):
    """the two ends of the chain side by side (text figure, thin rules only)"""
    rows = [("", "output end: before the lick", "input end: stimulus onset"),
            ("time", "100 ms before the first lick", "5-35 ms after the whisker"),
            ("axis", "rewarded - unrewarded lick (AH - SL)", "lick - no lick (active hit - miss)"),
            ("what moves", "active whisker hits, over session time", "passive whisker response, pre -> post"),
            ("drift reference", "spontaneous licks (same drift)", "auditory response; linear-shift null"),
            ("R+", "toward the rewarded lick", "no change"),
            ("R-", "away from the rewarded lick", "toward no-lick"),
            ("R+ vs R-", R["chain_pre"], R["chain_on"])]
    fig, ax = plt.subplots(figsize=(5.6, 2.3)); ax.axis("off"); ax.set_xlim(0, 1); ax.set_ylim(0, len(rows))
    xs = (0.0, 0.2, 0.61)
    for i, r in enumerate(rows):
        y = len(rows) - i - 0.5
        for j, t in enumerate(r):
            col = COL[r[0]] if r[0] in COL and j > 0 else "black"
            ax.text(xs[j], y, t, va="center", fontsize=6.2, weight="bold" if i == 0 or j == 0 else "normal", color=col)
        if i in (0, len(rows) - 2):
            ax.axhline(len(rows) - i - 1, color="black", lw=0.6)
    return D.savefig(fig, "story_chain_table")


# ----------------------------------------------------------------------------------------------------------- main
def main():
    H.setup()
    plt.rcParams.update({"mathtext.fontset": "dejavusans"})
    R = {}
    N = json.loads((D.RDIR / "numbers.json").read_text())
    A = D.load_135()
    C = pd.read_parquet(EA / "140_coding_direction_noise.parquet"); C = C[(C.split == "hitmedian") & C.skipped_reason.isna()]
    V = pd.read_parquet(EA / "146_choice_axis_readout.parquet")
    V = V[(V.area == "whole_brain") & (V.unit_set == "stable") & V.skipped_reason.isna()]
    VE, VB, VT = (V[V.response == r] for r in ("epochbase", "baseline", "trialbase"))
    L = H.learners(); AL = A[H.mouse_of(A).isin(L)]
    figs = dict(task=D.fig_task(), timeline=D.fig_timeline(), splits=D.fig_splits(R), cd=fig_cd_schematic(), wd=D.fig_wd(R),
                bridge=D.fig_bridge(), pa=D.fig_pattern_axis(), epochs=D.fig_lick_epochs(R), dtrain=D.fig_dec_train(),
                dro=D.fig_dec_readout(R, VE), dnull=D.fig_dec_null(R, VE, VB), null=D.fig_null(),
                b4=D.excess_fig("b4", [(A, "shift_excess_dWR", "whisker\ncosine"), (A, "shift_excess_dAR", "auditory\ncosine"),
                                       (A, "shift_excess_dWAP", "whisker - auditory\nprojection"),
                                       (VT, "shift_excess_dWA", "decoder: whisker -\nauditory readout")], R, w=6.9),
                b5=D.excess_fig("b5", [(AL, "shift_excess_dWR", "learners:\nwhisker cosine"), (AL, "shift_excess_dWAP", "learners:\nW - A projection"),
                                       (C, "shift_excess_daxisR", "late CD: passive\naxis cosine"), (C, "shift_excess_dWAP", "late CD:\nW - A projection")], R, w=7.6),
                expo=fig_exposure_story(R), state=D.fig_state(R), syn=D.fig_synthesis())
    wd = R["wd"]; wc = {c: wd[(wd.panel == "c") & (wd.cohort == c)].iloc[0] for c in ("R+", "R-", "R+ vs R-")}
    b4 = R["b4"]; g4 = lambda m: b4[(b4.measure == m) & (b4.cohort == "R+ vs R-")].iloc[0]
    R["chain_pre"] = f"p = {P(wc['R+ vs R-'].p_nonparam, wc['R+ vs R-'].p_param)} (drift)"
    R["chain_on"] = f"p = {P(g4('shift_excess_dWR').p_nonparam, g4('shift_excess_dWR').p_param)} (whisker cosine, excess)"
    figs["chain"] = fig_chain_table(R)
    E = dict(
        cd=D.eq("cd", [r"$c_i = \dfrac{\mathbf{x}_i\cdot\widehat{\mathbf{CD}} - \langle \cdot \rangle_{\mathrm{SL}}}{\langle \cdot \rangle_{\mathrm{AH}} - \langle \cdot \rangle_{\mathrm{SL}}}$,   $\widehat{\mathbf{CD}} \propto \bar{\mathbf{x}}_{\mathrm{AH}} - \bar{\mathbf{x}}_{\mathrm{SL}}$"]),
        drift=D.eq("drift", [r"$c_i = a_k + b_k\,\tau_i,\quad \beta = b_{\mathrm{WH}} - b_{\mathrm{SL}}$"]),
        lick=D.eq("lick", [r"$\mathbf{L} = \bar{\mathbf{z}}_{\mathrm{hit}} - \bar{\mathbf{z}}_{\mathrm{miss}},\quad \mathbf{p}_W = \langle \boldsymbol{\varepsilon}_t \rangle_{t \in W}$"]),
        cos=D.eq("cos", [r"$\cos\theta = \dfrac{\mathbf{p}\cdot\mathbf{L}}{\|\mathbf{p}\|\,\|\mathbf{L}\|},\quad \pi = \dfrac{\mathbf{p}\cdot\hat{\mathbf{L}}}{\sqrt{n}}$"]),
        ro=D.eq("ro", [r"$\mathrm{readout} = \dfrac{s - \frac{1}{2}(\bar{s}_{\mathrm{hit}} + \bar{s}_{\mathrm{miss}})}{\mathrm{SD}(s_{\mathrm{active}})}$"]),
        excess=D.eq("excess", [r"$\Delta M = M^{\mathrm{post}} - M^{\mathrm{pre}}$",
                               r"$\mathrm{excess} = \Delta M - \dfrac{1}{K}\sum_{k=1}^{K}\Delta M^{(k)}_{\mathrm{shift}}$"]),
        contrast=D.eq("contrast", [r"$\Delta\pi_W - \Delta\pi_A$  (drift shared by both stimuli cancels)"]),
        cd2=D.eq("cd2", [r"$\mathbf{CD}_2 = \bar{\mathbf{z}}_{\mathrm{hit},2} - \bar{\mathbf{z}}_{\mathrm{miss},2}$"]),
        ss=D.eq("ss", [r"$\hat{\mathbf{x}} = \mathbf{L}/\|\mathbf{L}\|,\quad \Delta x = (\mathbf{p}^{\mathrm{post}} - \mathbf{p}^{\mathrm{pre}})\cdot\hat{\mathbf{x}}$"]),
        split=D.eq("split", [r"midpoint: $t_{n/2}$;   hit-median: hit $\lfloor H/2 \rfloor + 1$"]),
        expo=D.eq("expo", [r"$\Delta M = \beta_0 + \beta_{R-}\,\mathbb{1}_{R-} + \beta_W z(n_W) + \beta_A z(n_A) + \epsilon$"]))
    ns, bl, tk, eqb, pic = D.new_slide, D.bullets, D.takehome, D.eqbox, D.picture
    m_ = lambda d, c, col: float(np.nanmean(d[d.reward_group == c][col]))

    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    D.CUR.update(sec="intro", n=1)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    D.text(s, 0.7, 0.8, 12, 1.6, "Within one learning session, the reward contingency re-maps the whisker sensorimotor chain at both ends",
           size=32, bold=True)
    D.text(s, 0.7, 2.6, 12, 0.5, "Whole-brain Neuropixels, single-session whisker learning; R+ (whisker licks rewarded) vs R- (not rewarded)", size=16)
    for i, (k, t) in enumerate((("1", "Output end: just before the lick, whisker hits move toward or away from rewarded licks"),
                                ("2", "Input end: the earliest whisker response moves away from licking in R-, beyond session time"),
                                ("3", "Both ends together; controls; caveats"))):
        D.text(s, 0.7, 3.9 + i * 1.0, 0.75, 0.6, k, size=24, bold=True)
        D.text(s, 1.4, 3.95 + i * 1.0, 11.3, 0.8, t, size=18)
    # task
    s = ns(prs, "Same stimuli, same action, different contingency", "Learning session (whisker day 0), one per mouse")
    pic(s, figs["task"], 0.5, 1.4, w=6.2); pic(s, figs["timeline"], 0.5, 5.3, w=6.2)
    bl(s, 7.2, 1.5, 5.7, 4.5, ["Hit = lick after the whisker stimulus: same label in both cohorts (trained response in R+, error in R-).",
                               "Question: does the contingency change how the whisker is represented relative to licking and reward, "
                               "within the session in which the mice learn it?",
                               "Statistics: session = unit (one per mouse); a non-parametric and a parametric test (p | p); uncorrected."], size=15)
    # decoding is stable
    sm, sh, hr = R["split_mid"], R["split_hitmedian"], R["split_hr"]
    k0 = lambda k: N.get(k, np.nan)
    s = ns(prs, "Choice decoding is stable within the session, in both cohorts",
           "Hit vs miss decoded from 5-100 ms whole-brain activity (above a linear-shift null), before vs after a split of the session")
    pic(s, figs["splits"], 0.3, 1.3, w=8.6)
    eqb(s, 9.2, 1.3, 3.8, [("split points", E["split"])])
    bl(s, 9.2, 2.6, 3.8, 3.0, [f"Session-wide (5-50 ms): R+ {k0('k0_rp'):+.3f}, R- {k0('k0_rm'):+.3f} above the null; R+ vs R- p = "
                               f"{H.pnum(k0('k0_mw'))} | {H.pnum(k0('k0_w'))}.",
                               f"Change between halves, R+ vs R-: midpoint p = {P(*sm['cohort'])}, hit-median p = {P(*sh['cohort'])}.",
                               "Exception: a small R+ gain at the behavioural change point, in abrupt learners (backup)."], size=12)
    tk(s, 9.2, 5.6, 3.8, 1.3, "Whether a lick follows is equally decodable throughout; the question is what the code looks like.")
    # 1 output end
    D.section_slide(prs, "1", "Output end: just before the lick", "Where do whisker hits sit between a rewarded and an unrewarded lick? (within-day project)")
    s = ns(prs, "A reward-lick axis separates rewarded from unrewarded licks", "Each session's coding direction from auditory hits (always rewarded) to spontaneous licks (never rewarded)")
    pic(s, figs["cd"], 0.4, 1.4, w=7.0)
    eqb(s, 7.8, 1.3, 5.2, [("projection of trial i (SL = 0, AH = 1)", E["cd"]), ("per-session drift over normalised session time", E["drift"])])
    bl(s, 0.5, 4.8, 7.0, 2.0, ["Both ends of the axis are licks: the movement is matched, the expected outcome is not.",
                               "Whisker hits are placed on the axis trial by trial; their drift is measured relative to the spontaneous "
                               "licks', so drift shared by all licks (time, state) cancels."], size=13)
    tk(s, 7.8, 4.8, 5.2, 1.5, "The axis asks: does the brain treat a whisker lick like a lick that will be rewarded?")
    s = ns(prs, "Before the lick, R+ whisker hits converge on rewarded licks, R- whisker hits diverge",
           "100 ms before the first lick; learning-session curves (all day-0 sessions) and all expert sessions (unpaired); >= 4 whisker hits")
    pic(s, figs["wd"], 0.3, 1.3, w=8.6)
    bl(s, 9.2, 1.4, 3.8, 3.4, [f"Drift R+ {wc['R+'].mean_a:+.2f} (p = {P(wc['R+'].p_nonparam, wc['R+'].p_param)}), R- {wc['R-'].mean_a:+.2f} "
                               f"(p = {P(wc['R-'].p_nonparam, wc['R-'].p_param)}).",
                               f"R+ vs R- p = {P(wc['R+ vs R-'].p_nonparam, wc['R+ vs R-'].p_param)}; {wc['R+ vs R-'].note}.",
                               "Expert sessions: R+ high on the axis, R- low (dots, right)."], size=12)
    tk(s, 9.2, 5.0, 3.8, 1.9, "Before the lick, R+ whisker hits come to look like rewarded licks and R- whisker hits less so, "
       "while decoding accuracy does not change.")
    # 2 input end
    D.section_slide(prs, "2", "Input end: the earliest whisker response", "Does the 5-35 ms whisker response, measured without the task, move relative to licking?")
    s = ns(prs, "From the response to its direction", "Passive trials before and after the task; the active trials give a reference for licking")
    pic(s, figs["bridge"], 0.4, 1.3, w=7.2)
    bl(s, 0.5, 4.75, 7.1, 2.3, ["A change in gain scales the whisker response along its own direction.",
                                "The active trials define the direction of licking: the lick axis L = whisker hits - whisker misses.",
                                "Passive trials are measured without the task, so a move along L is a change in the sensory "
                                "representation itself."], size=13)
    tk(s, 7.9, 1.6, 5.1, 1.6, "After the task, does the whisker response look more or less like a lick trial?")
    s = ns(prs, "One pattern, one axis", "Passive whisker pattern p (before or after the task) against the active lick axis L")
    pic(s, figs["pa"], 0.6, 1.4, w=4.6); pic(s, figs["timeline"], 0.4, 5.4, w=5.6)
    eqb(s, 6.4, 1.3, 6.6, [("lick axis (active whisker trials) and passive whisker pattern", E["lick"]), ("direction and reach", E["cos"])])
    bl(s, 6.4, 4.2, 6.6, 2.6, ["Same tracked, drift-checked units in every epoch; active trials with a lick before 35 ms excluded.",
                               "Prediction: after the task, the R- whisker pattern points less along L; R+ unchanged; auditory unchanged."], size=13)
    d = R["d"].iloc[0]
    s = ns(prs, "After the task, the R- whisker response leaves the lick axis", "Cosine of the evoked patterns with the lick axis, in the four epochs")
    pic(s, figs["epochs"], 0.6, 1.3, w=7.6)
    bl(s, 8.6, 1.5, 4.4, 3.4, [f"Passive post - pre, R+ vs R-: p = {P(d.p_nonparam, d.p_param)}.",
                               "Auditory pattern (dashed): no cohort difference.",
                               "Active halves are descriptive (they average hits and misses).",
                               "Raw change: a session-time check follows."], size=14)
    tk(s, 8.6, 5.2, 4.4, 1.5, "R- whisker responses decouple from the lick axis during the task and stay decoupled.")
    dro, droA = R["dro_W"], R["dro_A"]
    s = ns(prs, "A choice decoder reads the R- whisker response as a no-lick trial",
           "Decoder trained on active hit vs miss (5-35 ms), applied to passive trials it never saw")
    pic(s, figs["dtrain"], 0.4, 1.35, w=4.6); pic(s, figs["dro"], 5.6, 1.3, w=7.4, h=3.3)
    eqb(s, 0.5, 4.2, 4.5, [("score s of a trial; > 0 hit side, < 0 miss side (SD of held-out active scores)", E["ro"])])
    bl(s, 5.6, 4.7, 7.4, 1.1, [
        f"Whisker trials: R- {m_(VE, 'R-', 'ro_std_passive_pre_W'):+.2f} -> {m_(VE, 'R-', 'ro_std_passive_post_W'):+.2f} SD "
        f"(p = {P(dro['R-'][1], dro['R-'][2])}); R+ {m_(VE, 'R+', 'ro_std_passive_pre_W'):+.2f} -> {m_(VE, 'R+', 'ro_std_passive_post_W'):+.2f} "
        f"(p = {P(dro['R+'][1], dro['R+'][2])}); R+ vs R- p = {P(*dro['cohort'])}.",
        f"Auditory trials: R+ vs R- p = {P(*droA['cohort'])}."], size=12)
    tk(s, 0.5, 5.85, 12.4, 0.95, "A different readout, the same direction: after the task, R- whisker responses read as no-lick trials.")
    dbs, dcov = R["dbase"], R["dcov"]
    VEd = VE.assign(d=VE.ro_std_passive_post_W - VE.ro_std_passive_pre_W)
    s = ns(prs, "But a decoder trained on time-shifted labels does it too",
           "Shifted labels keep the session's drift and destroy the trial-by-trial hit / miss match; the baseline window carries no sensory response")
    pic(s, figs["dnull"], 0.3, 1.3, w=8.9)
    bl(s, 9.4, 1.4, 3.6, 3.9, [
        f"R- change: real {m_(VEd, 'R-', 'd'):+.2f} SD, shifted-label decoders {m_(VE, 'R-', 'shift_null_dW_mean'):+.2f}.",
        f"Baseline window alone: R- change {dbs['R-'][0]:+.2f} (p = {P(dbs['R-'][1], dbs['R-'][2])}); R+ vs R- p = {P(*dbs['cohort'])}.",
        f"R- coefficient: p = {H.pnum(dcov[0][2])} alone, {H.pnum(dcov[1][2])} with baseline rate and gap, {H.pnum(dcov[2][2])} with reward rate."],
        size=12)
    tk(s, 9.4, 5.3, 3.6, 1.6, "R- hits are early and misses late: any hit / miss axis can read session time and state.")
    s = ns(prs, "A null that keeps the drift", "Rebuild every hit / miss axis from labels shifted against the trials; compare the real change with theirs")
    pic(s, figs["null"], 0.5, 1.5, w=7.8)
    eqb(s, 8.7, 1.3, 4.3, [("change of a metric, and excess over K = 50 axes rebuilt from shifted labels", E["excess"]),
                           ("whisker - auditory contrast", E["contrast"])])
    bl(s, 0.6, 5.0, 12, 1.9, ["Shuffled labels destroy trial order: an axis that learned 'late = miss' is not in that null.",
                              "The linear shift keeps the drift of labels and activity; conservative, since real learning is time-correlated too.",
                              "The auditory response shares the session's drift and state but not the whisker contingency."], size=14)
    s = ns(prs, "Beyond session time, the R- whisker response still moves away from licking",
           "Passive pre -> post change minus the change produced by axes rebuilt from shifted labels")
    pic(s, figs["b4"], 0.3, 1.3, w=8.6)
    bl(s, 9.2, 1.4, 3.8, 3.8, [f"Lick axis, whisker cosine: R+ vs R- p = {P(g4('shift_excess_dWR').p_nonparam, g4('shift_excess_dWR').p_param)}.",
                               f"Auditory cosine p = {P(g4('shift_excess_dAR').p_nonparam, g4('shift_excess_dAR').p_param)}.",
                               f"Whisker - auditory projection p = {P(g4('shift_excess_dWAP').p_nonparam, g4('shift_excess_dWAP').p_param)}.",
                               f"Decoder: whisker alone p = {P(*H.unpaired(VE[VE.reward_group == 'R+'].shift_excess_dW, VE[VE.reward_group == 'R-'].shift_excess_dW))}; "
                               f"whisker - auditory p = {P(g4('shift_excess_dWA').p_nonparam, g4('shift_excess_dWA').p_param)}."], size=12)
    tk(s, 9.2, 5.3, 3.8, 1.6, "The lick axis shows the R- whisker decoupling directly; the decoder only relative to auditory.")
    b5 = R["b5"]; g5 = lambda i: b5[b5.cohort == "R+ vs R-"].iloc[i]
    s = ns(prs, "Robust in learners only and with a second hit / miss axis", "Learners only (lick axis); late coding direction of a hit-median split")
    pic(s, figs["b5"], 0.3, 1.3, w=9.4)
    eqb(s, 9.9, 1.3, 3.1, [("late coding direction", E["cd2"])])
    bl(s, 9.9, 2.4, 3.1, 3.0, [f"Learners: p = {P(g5(0).p_nonparam, g5(0).p_param)} and {P(g5(1).p_nonparam, g5(1).p_param)}.",
                               f"Late CD: p = {P(g5(2).p_nonparam, g5(2).p_param)} and {P(g5(3).p_nonparam, g5(3).p_param)}."], size=12)
    tk(s, 9.9, 5.4, 3.1, 1.5, "Same answer in learners and with a different hit / miss axis.")
    eo, ec = R["expo_ols"], R["expo_corr"]
    gw = lambda y, mdl: eo[(eo.measure == y) & (eo.model == mdl)].iloc[0]
    cw = {c: ec[(ec.measure == "lick_WR") & (ec.cohort == c) & (ec.x == "nW")].iloc[0] for c in COH}
    s = ns(prs, "Not explained by how many stimuli the mice received", "R- sessions are longer: more whisker and auditory trials, in the same proportion")
    pic(s, figs["expo"], 0.3, 1.3, w=9.0)
    eqb(s, 9.5, 1.3, 3.5, [("covariate model", E["expo"])])
    bl(s, 9.5, 2.45, 3.5, 3.0, [f"Whisker trials R+ vs R- p = {P(*R['expo_n_nW'])}; whisker fraction {R['expo_frac'][0]:.2f} vs {R['expo_frac'][1]:.2f}.",
                                f"Within cohorts, lick-axis excess vs whisker trials: R+ p = {P(cw['R+'].p_nonparam, cw['R+'].p_param)}, "
                                f"R- p = {P(cw['R-'].p_nonparam, cw['R-'].p_param)}.",
                                f"With both counts as covariates: cohort p = {H.pnum(gw('lick_WR', '+ counts').p)} (whisker cosine), "
                                f"{H.pnum(gw('lick_WAP', '+ counts').p)} (W - A projection)."], size=12)
    tk(s, 9.5, 5.4, 3.5, 1.5, "The auditory tone is also presented more often in R-, and it does not move.")
    f_ = R["f"].iloc[0]
    s = ns(prs, "Geometry: the R- whisker response slides toward the miss level", "Passive pre at the origin; x = lick axis; dashed: where active misses sit on x")
    pic(s, figs["state"], 0.8, 1.3, w=6.8)
    eqb(s, 8.0, 1.3, 5.0, [("x axis and displacement", E["ss"])])
    bl(s, 8.0, 2.6, 5.0, 2.6, [f"Displacement along x, R+ vs R-: p = {P(f_.p_nonparam, f_.p_param)} (raw, no session-time null: an illustration of "
                               "the result above, not a separate test).",
                               "y = the response's own pre-task pattern; its pre point is built from the same trials (not cross-validated), "
                               "so the downward shift along y is not interpreted."], size=12)
    tk(s, 8.0, 5.4, 5.0, 1.4, "After the task, an R- whisker stimulus evokes a pattern like a whisker trial without a lick.")
    # 3 both ends
    D.section_slide(prs, "3", "Both ends", "The same question at the output and the input of the whisker sensorimotor chain")
    s = ns(prs, "Two ends of the chain, one logic", "Position of whisker-driven activity on a behaviourally defined axis, against a reference that shares the drift")
    pic(s, figs["chain"], 0.5, 1.3, w=8.4)
    bl(s, 9.2, 1.4, 3.8, 3.6, ["Spontaneous licks are to the pre-lick analysis what the auditory response is to stimulus onset: same "
                               "session, same drift, different contingency.",
                               "R+: the whisker becomes a reward lick at the output end; the early sensory response keeps its relation to licking.",
                               "R-: the whisker loses its link to reward and licking at both ends."], size=12)
    tk(s, 9.2, 5.2, 3.8, 1.6, "R+ re-maps at the output end, R- at both ends.")
    s = ns(prs, "Synthesis", "Decodability is stable; the representation relative to reward and licking is re-mapped, in opposite directions")
    pic(s, figs["syn"], 0.6, 1.4, w=8.4)
    bl(s, 9.2, 1.5, 3.8, 4.0, ["R+: pre-lick whisker hits converge on the reward-lick representation within the learning session.",
                               "R-: pre-lick whisker hits diverge from it, and the earliest whisker response decouples from the lick axis.",
                               "Neither is visible in decoding accuracy."], size=14)
    tk(s, 0.6, 5.7, 12.2, 1.1, "The reward contingency re-maps the whisker sensorimotor chain within one session; decodability alone misses it.")
    figs["expert"] = fig_expert_story(R)
    T = R["expert"]
    gx = lambda m, test, grp: T[(T.measure == m) & (T.test == test) & (T.group == grp)].iloc[0]
    Px = lambda r: P(r.p_nonparam, r.p_param)
    nx = T[(T.measure == "shift_excess_dWR") & (T.test == "vs 0 (Wilcoxon | t)")].set_index("group").n
    s = ns(prs, "Expert sessions as a control", "Same pipeline on expert days; filled: learning (day 0), open: expert; dots = sessions (156)")
    pic(s, figs["expert"], 0.3, 1.3, w=12.6, h=3.9, center=True)
    bl(s, 0.5, 5.25, 7.6, 1.8, [
        "R+ vs R- on expert days: " + "; ".join(f"{lab} p = {Px(gx(m, 'R+ vs R- (Mann-Whitney | Welch)', 'expert'))}" for m, lab in
                                               (("shift_excess_dWR", "whisker cosine"), ("shift_excess_dWAP", "W - A projection"),
                                                ("shift_excess_dWA", "decoder W - A"))) + ".",
        "Cohort x stage: " + "; ".join(f"{lab} p = {Px(gx(m, 'cohort x stage (permutation | OLS)', 'interaction'))}" for m, lab in
                                       (("shift_excess_dWR", "whisker cosine"), ("shift_excess_dWAP", "W - A projection"),
                                        ("shift_excess_dWA", "decoder W - A"))) + ".",
        f"Expert lick-axis sessions: R+ {int(nx['R+ expert'])}, R- {int(nx['R- expert'])}: expert R- mice rarely lick to the whisker."], size=12)
    tk(s, 8.4, 5.25, 4.6, 1.6, "No cohort difference within expert sessions; the contrast measures look learning-specific, "
       "but the R- expert group is small.")
    s = ns(prs, "Caveats")
    bl(s, 0.6, 1.4, 12.2, 5.4, [
        "Session time: hit / miss labels drift; stimulus-onset inference is on the excess over a conservative linear-shift null.",
        "State: the pre-stimulus window moves miss-ward after the task in R- (both stimuli); the whisker-specific claim rests on the "
        "whisker - auditory contrasts.",
        "Reward: R- mice collect fewer rewards (by design); reward rate and contingency cannot be separated.",
        "Pre-lick: a mixed model with the spontaneous-lick reference is weaker than the per-session test shown (within-day report).",
        "Units differ between the pre-lick (good + MUA) and stimulus-onset (tracked stable) analyses: the parallel is conceptual.",
        "Expert R- sessions rarely have enough whisker hits for a hit / miss axis.",
        "Exploratory, uncorrected."], size=16)
    # backup
    D.CUR["sec"] = "bk"
    PUB = D.PUB
    for path, t in ((PUB / "151_area_heatmap_all.png", "Area groups: post - pre changes (151)"),
                    (PUB / "138_cosyne_lt_placebo_all.png", "Learning trial vs placebo splits (138)"),
                    (PUB / "143_lt_split_L5_windows_all.png", "L5 change point across windows (143)"),
                    (PUB / "143_step_vs_gradual_all.png", "Step vs gradual learners (143)"),
                    (PUB / "144_hitmedian_split_all.png", "Hit-median vs midpoint split, all windows (144)"),
                    (PUB / "150_axis_alignment_shift_null_all.png", "Lick axis and late coding direction vs the shift null (150)"),
                    (PUB / "149_passive_readout_shift_null_all.png", "Decoder readout vs the shift null (149)"),
                    (PUB / "147_controls.png", "State and engagement controls (147)"),
                    (PUB / "153_exposure_control.png", "Exposure control, full figure (153)"),
                    (PUB / "154_cosyne_remapping_all.png", "COSYNE figure (154)")):
        if path.exists():
            s = ns(prs, t); pic(s, path, 0.4, 1.2, w=12.5, h=5.8, scale=20, center=True)
    prs.save(OUT)
    print(OUT, D.CUR["n"], "slides")


if __name__ == "__main__":
    main()
