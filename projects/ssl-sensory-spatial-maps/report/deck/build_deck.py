"""Sensory maps + stimulus-arrival slide deck (16:9 PowerPoint) from the panels in img/ (made on haas by
../build_deck_assets.py) and the numbers in report_numbers.json (written by ../build_report.py from the result tables;
nothing is hard-coded). Each slide: section label, claim title, figure(s) fitted to the free area, optional bullets,
footnote with n / test / effect size, speaker notes. 90 % projection zones are the main version (user 2026-10-04).
Run: python build_deck.py (needs python-pptx and Pillow).
"""
import json
import pathlib

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

HERE = pathlib.Path(__file__).parent
IMG = HERE / "img"
W, H = 13.333, 7.5
M = 0.4                                   # side margin (in)
TOP, FOOT = 1.15, 0.75                    # title band, footnote band
GREY = RGBColor(0x55, 0x55, 0x55)
LAT_GROUPS = ["SS-whisker", "Posterior parietal", "Auditory", "Midbrain", "Thalamus", "Motor", "Striatum"]


def pt(p):
    return "p < 0.001" if p < 0.001 else f"p = {p:.3f}" if p < 0.01 else f"p = {p:.2f}"


def onset_text(lst, k=None):
    q = [(a, v) for a, v in lst if v is not None][:k]
    return ", ".join(f"{a} {v:.0f}" for a, v in q)


def slides(N):
    A, Z, C9, C7 = N["arrival"], N["zones"], N["coloc90"], N["coloc70"]
    lat = {g: (w, a) for g, w, a in N["latency_groups"]}
    g_on = [(a, v) for a, v in A["groups"] if v is not None]
    f_on = [(a, v) for a, v in A["areas"] if v is not None]
    lo, hi = g_on[0][1], g_on[-1][1]
    first = [a for a, v in g_on if v == lo]
    first_f = [a for a, v in f_on if v == f_on[0][1]] if f_on else []
    rw, ra = N["resp"]["whisker"], N["resp"]["auditory"]
    mod = N["resp"]["modality"]
    resp = round((rw[0] + ra[0]) / 2)
    matched = {a: v for a, v in A["matched"]}
    mtxt = ", ".join(f"~{matched[a]:.0f} {a.lower()}" if matched.get(a) is not None else f"> 2,000 {a.lower()}"
                     for a in ("Midbrain", "Striatum", "Motor") if a in matched)
    prelim = f"{len(g_on)} area groups and {len(f_on)} fine areas with an onset"
    srcs = ", ".join(f"{s} {n} experiments / {l} lines" for s, n, l in Z["sources"])
    mo = N["modality_offset"]
    lk = N.get("link") or {}
    pas = A.get("passive_groups")
    S = [
        ("", "Conclusions", [], [
            f"About {resp} % of neurons respond to each stimulus, in different places; modality preference is spatially organised",
            f"Each stimulus reaches its own sensory system first: whisker → whisker cortex (~{lat['SS-whisker'][0]:.0f} ms), "
            f"sound → midbrain (~{lat['Midbrain'][1]:.0f} ms)",
            f"Population decoding detects stimulus modality within {lo:.0f}–{hi:.0f} ms: first {', '.join(first)}; last {g_on[-1][0]}",
            f"Faster areas also carry more early information (ρ = {A['rho_groups']:.2f})",
            f"Whisker- and auditory-cortex projections overlap in {Z['overlap90']:.0f} mm³; bimodal neurons are enriched there "
            f"({C9['diff']:+.1f} points); within shared targets the two preferences are spatially offset"],
         f"Generated {N['generated']} from the result tables; decoding: {prelim}.", "Sensory coding first (where, when, population), then convergence."),
        ("Data", f"{N['n_neurons']:,} neurons from {N['n_sessions']} sessions, all registered to the Allen CCF", [("flat_recorded.png", 0.38)], [
            f"{N['n_mice']} mice, R+ and R− cohorts, learning day and expert days (pooled)",
            "Good + multi-unit clusters (Kilosort 4)",
            "Probe-track reconstruction, folded onto the left hemisphere",
            f"{N['n_structures']} recorded structures (≥ 10 neurons)",
            "Whisker stimulus artefact (−10 to +5 ms) corrected on all spike trains"],
         "Top: recorded isocortex neurons on the Allen butterfly flatmap (anterior up). Bottom: recorded-neuron density (σ = 150 µm).",
         "Active trials: perf ≠ 6, disengaged tail trimmed. Passive: fixed ~3 s ITI blocks."),
        ("I · Where neurons respond", f"About {resp} % of neurons respond to each stimulus, in different places", [("flat_whisker_auditory.png", 1)], [],
         f"Whisker: {rw[0]:.1f} % significant ({rw[1]:.1f} % excited); auditory: {ra[0]:.1f} % ({ra[1]:.1f} % excited). ROC 5–35 ms vs "
         "pre-trial baseline, 1000 permutations, p < 0.05. Bottom: selectivity density normalised by recorded-neuron density (σ = 150 µm).",
         "Selectivity = 2·AUC − 1. Flatmap: left hemisphere, anterior up."),
        ("I · Where neurons respond", "Modality preference is spatially organised", [("flat_modality.png", 0.32)], [
            f"{mod[0]:.1f} % of neurons prefer one modality",
            f"Auditory-preferring {mod[1]:.1f} %, whisker-preferring {mod[2]:.1f} %",
            "Whisker-preferring neurons cluster in somatosensory cortex; auditory-preferring neurons are more widespread"],
         "ROC whisker vs auditory trials, 5–35 ms, active task, permutation p < 0.05. Yellow: whisker-preferring; blue: auditory-preferring.", ""),
        ("II · When neurons respond", "Each modality reaches its own sensory system first", [("flat_latency.png", 0.55)], "TABLE",
         f"Median half-time to peak of responsive neurons (light = fast); all neurons: whisker {N['latency_median']['whisker']:.1f} ms, "
         f"auditory {N['latency_median']['auditory']:.1f} ms. Whisker latency searched after +5 ms (artefact window).",
         "Latency = last upward crossing of half the peak before the peak; 1 ms PSTH, σ = 2 ms."),
        ("III · Population coding", f"Stimulus modality is decodable within {lo:.0f}–{hi:.0f} ms, first in {' and '.join(first).lower()}",
         [("arr_groups_heatmap_ranking.png", 1)], [],
         "Pseudo-populations of 200 neurons (20 sessions × 10), whisker vs auditory task trials, L2 logistic regression per 20 ms bin. "
         f"Corrected accuracy = real − label-shuffled. Grey: not above chance. Onset: ≥ 80 % of the next 25 ms above chance. {prelim}.",
         "Onsets (ms): " + onset_text(A["groups"])),
        ("III · Population coding", f"Fine areas: {', '.join(first_f)} first", [("arr_areas_heatmap_ranking.png", 1)], [],
         "Onsets (ms): " + onset_text(A["areas"], 12) + ("…" if len(f_on) > 12 else ""), ""),
        ("III · Population coding", "Single-neuron latency and population onset", [("link_latency_onset.png", 1)], [],
         "Population decoding onset (N = 200) vs median single-neuron latency per area group; OLS line with 95 % CI. "
         + (f"Faster modality: ρ = {lk['rho']:.2f}, {pt(lk['p'])}, n = {lk['n']} area groups." if lk else ""), ""),
        ("III · Population coding", "Faster areas carry more early information", [("arr_groups_controls.png", 1)], [],
         f"Early accuracy (5–50 ms) grows with neuron count; onset vs early accuracy: Spearman ρ = {A['rho_groups']:.2f}. "
         f"Matching whisker cortex at 100 neurons needs {mtxt} neurons.",
         "Matched-accuracy control: neuron count at which each group reaches the early accuracy of SS-whisker at 100 neurons."),
        ("III · Population coding",
         f"Passive trials (no task, no licks): onsets {pas[0][1]:.0f}–{[v for _, v in pas if v is not None][-1]:.0f} ms" if pas else
         "Passive trials (no task, no licks): pending", [("pas_groups_heatmap_ranking.png", 1)],
         [] if pas else ["Passive whisker vs auditory decoding (sessions with ≥ 3 passive trials per stimulus)",
                         "Tests whether the onset order is sensory rather than task- or lick-related"],
         ("Onsets (ms): " + onset_text(pas)) if pas else "Shown here once the passive sweep finishes.", ""),
        ("IV · Convergence", "Whisker and auditory cortex project to partly overlapping targets", [("projection_zones_coronal.png", 1)], [],
         f"Allen anterograde tracing, all projection-neuron lines averaged per line ({srcs}); 90 % projection zones. Overlap "
         f"{Z['overlap90']:.1f} mm³ (70 % zones: {Z['overlap70']:.1f} mm³); largest recorded pieces: {', '.join(Z['pieces'])}.",
         "Question for this section: where do the two streams described so far meet?"),
        ("IV · Convergence", "Cortical responses against the projection zones", [("flat_whisker_auditory_zones.png", 0.62), ("flat_bimodal_zones.png", 0.33)], [],
         "Same flatmaps as before with the 90 % projection zones of whisker (teal) and auditory (brown dashed) cortex. Right: bimodal "
         "neurons (responsive to both stimuli, Bonferroni over the session's active / passive tests).", ""),
        ("IV · Convergence", "Bimodal neurons are enriched where the projections converge", [("colocation_figure.png", 0.58)], [
            "Bimodal = responsive to both stimuli",
            f"{N['bimodal_pct']:.0f} % of responsive neurons are bimodal",
            f"Inside the overlap: {C9['P_in']:.1f} % vs {C9['P_ref']:.1f} % overall",
            f"{C9['diff']:+.1f} points (95 % CI {C9['ci_lo']:+.1f} to {C9['ci_hi']:+.1f}), {pt(C9['p'])}",
            f"Stricter 70 % zones: {C7['diff']:+.1f} points, {pt(C7['p'])}"],
         "Hierarchical bootstrap (sessions, then neurons) of P(bimodal | inside overlap) − P(bimodal | all responsive); 90 % zones. "
         "Co-location, not causation.", ""),
        ("IV · Convergence", "Within shared targets, the two preferences are spatially offset", [("modality_contours.png", 1)], [],
         f"{mo['n_sig']} of {mo['n']} target slabs"
         + (f": centroids {mo['dmin']:.0f}–{mo['dmax']:.0f} µm apart" if mo['dmin'] is not None else "")
         + " (within-session label permutation, Holm p < 0.05).",
         "80 % highest-density contours (Gaussian KDE). Permuting within sessions removes offsets caused by which sessions recorded which preference."),
        ("Summary", "Summary: separate entry routes, shared targets", [("summary_schematic.png", 1)], [],
         "Draft. a) Whisker route via thalamus to whisker cortex, sound via midbrain and thalamus to auditory cortex; both cortices "
         "project to the shared targets, where bimodal neurons are enriched. Times: population decoding onset (N = 200). b) Onsets "
         "with 95 % bootstrap CI.", "Routes: textbook anatomy plus the Allen projection zones; numbers read from the result tables."),
        ("", "Caveats and next steps", [], [
            "Co-location ≠ causation: tracing from other mice, axons of passage, ~100–200 µm CCF uncertainty",
            "The bimodal enrichment depends on the zone level (90 % main; weaker with 70 %)",
            "Task-trial decoding after ~100 ms can use licks → passive trials",
            "Whisker information cannot appear before +5 ms (artefact window)",
            "Pilot sampling (100 iterations × 10 shuffles); final N = 200 onsets with 1000 × 20"], "", ""),
        ("", "Backup: decoding method", [("arr_groups_method.png", 0.4)], [
            "20 sessions with replacement, N/20 neurons per session",
            "Pseudo-trials per class by balanced reuse of each session's trials",
            "L2 logistic regression per bin, 3-fold CV, inner 2-fold for C",
            "d(t) = BA_real − mean of 10 within-session label shuffles",
            "Above chance: 5th percentile of d over 100 iterations > 0",
            "Onset: first bin above chance with ≥ 80 % of the next 25 ms above chance"], "", ""),
        ("", "Backup: trial-shuffle control", [("arr_groups_shuffle.png", 0.5), ("arr_groups_timecourse.png", 0.5)], [], "", ""),
        ("", "Backup: all flatmaps with projection zones", [("flat_all_zones.png", 1)], [], "", ""),
    ]
    rows = [r for r in N["latency_groups"] if r[0] in LAT_GROUPS]
    bw = min(rows, key=lambda r: r[1])[0]
    ba = min(rows, key=lambda r: r[2])[0]
    table_rows = [("Area group", "Whisker (ms)", "Auditory (ms)", False, False)] + [
        (g, f"{w:.1f}", f"{a:.1f}", g == bw, g == ba) for g, w, a in rows]
    return S, table_rows


def text(slide, x, y, w, h, s, size, bold=False, color=None, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = s
    p.font.size, p.font.bold, p.font.italic = Pt(size), bold, italic
    if color is not None:
        p.font.color.rgb = color
    return tf


def bullets(slide, x, y, w, h, items, size):
    tf = text(slide, x, y, w, h, "• " + items[0], size)
    for it in items[1:]:
        p = tf.add_paragraph()
        p.text, p.font.size = "• " + it, Pt(size)
        p.space_before = Pt(8)


def picture(slide, f, x, y, w, h):
    """fit the image in the box, keep aspect, centre"""
    iw, ih = Image.open(f).size
    s = min(w / iw, h / ih)
    pw, ph = iw * s, ih * s
    slide.shapes.add_picture(str(f), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def table(slide, x, y, w, rows):
    t = slide.shapes.add_table(len(rows), 3, Inches(x), Inches(y), Inches(w), Inches(0.38 * len(rows))).table
    for i, (g, a, b, fw, fa) in enumerate(rows):
        for j, v in enumerate((g, a, b)):
            c = t.cell(i, j)
            c.text = v
            c.text_frame.paragraphs[0].font.size = Pt(14)
            c.text_frame.paragraphs[0].font.bold = (i == 0) or (j == 1 and fw) or (j == 2 and fa)


def main():
    N = json.loads((HERE / "report_numbers.json").read_text(encoding="utf-8"))
    S, rows = slides(N)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    text(s, 1, 1.9, W - 2, 1.6, "Whisker and auditory signals across the mouse brain", 40, bold=True)
    text(s, 1, 3.9, W - 2, 1, "Where, when and through which pathways: sensory maps and stimulus-arrival decoding", 22, color=GREY)
    text(s, 1, 4.9, W - 2, 0.6, f"SSL Neuropixels dataset, working version {N['generated'][:10]} (unpublished)", 16, color=GREY)

    missing = []
    for kick, title, imgs, bl, foot, notes in S:
        s = prs.slides.add_slide(blank)
        text(s, M, 0.3, W - 2 * M, 0.8, title, 28, bold=True)
        if kick:
            text(s, M, 0.02, 4, 0.3, kick.upper(), 10, bold=True, color=GREY)
        y0, y1 = TOP, H - (FOOT if foot else 0.3)
        imgs = [(f, fr) for f, fr in imgs if (IMG / f).exists() or missing.append(f)]
        fig_frac = sum(fr for _, fr in imgs)
        x = M
        for f, fr in imgs:
            w = (W - 2 * M) * fr
            picture(s, IMG / f, x, y0, w, y1 - y0)
            x += w
        rest = W - M - x
        if bl == "TABLE":
            table(s, x + 0.3, y0 + 0.6, rest - 0.6, rows)
        elif bl:
            if fig_frac == 0:
                bullets(s, 1.0, y0 + 0.2, W - 2.0, y1 - y0, bl, 22)
            else:
                bullets(s, x + 0.2, y0 + 0.4, max(rest - 0.4, 3), y1 - y0, bl, 16 if len(bl) > 4 else 18)
        if foot:
            text(s, M, H - FOOT, W - 2 * M, FOOT - 0.05, foot, 11, italic=True, color=GREY)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    out = HERE / "sensory_deck.pptx"
    prs.save(out)
    print(out, len(prs.slides), "slides; missing:", missing)


if __name__ == "__main__":
    main()
